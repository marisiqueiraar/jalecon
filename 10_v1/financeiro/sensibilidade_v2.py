"""Roda cenários e sensibilidade do modelo v2 recalculando no LibreOffice, e grava as abas
Cenarios e Sensibilidade (valores) no xlsx final.

Uso: python 10_v1/financeiro/sensibilidade_v2.py
"""
import json
import os
import subprocess
import tempfile

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

HERE = os.path.dirname(os.path.abspath(__file__))
BUILD = os.path.join(HERE, "build_model_v2.py")
FINAL = os.path.join(HERE, "Jalecon_Modelo_Financeiro_v2.xlsx")
WORK = tempfile.mkdtemp(prefix="jalecon_v2_")

OUTPUTS = [
    ("arr5", "Receita recorrente anualizada no ano 5"),
    ("arr7", "Receita recorrente anualizada no ano 7"),
    ("eq", "Mês de equilíbrio (EBITDA positivo daí em diante)"),
    ("cap", "Capital necessário (pior caixa acumulado)"),
    ("mb5", "Margem bruta no ano 5"),
    ("ltvcac", "LTV/CAC do plano principal (Consultório / Agenda Cheia), mídia paga"),
    ("payback", "Recuperação do CAC do plano principal (meses)"),
    ("m50", "Mês em que a receita anualizada chega a R$ 50 mi"),
    ("pen50", "Participação no mercado nesse mês"),
    ("rpp", "Receita por pessoa no ano 5"),
    ("cli5", "Clientes ativos no ano 5"),
]


def build(name, overrides):
    xlsx = os.path.join(WORK, name + ".xlsx")
    ovr = os.path.join(WORK, name + ".json")
    json.dump(overrides, open(ovr, "w"))
    subprocess.run(["python3", BUILD, xlsx, ovr], check=True, capture_output=True)
    return xlsx


def recalc_all(paths):
    out = os.path.join(WORK, "out")
    os.makedirs(out, exist_ok=True)
    for i in range(0, len(paths), 12):
        subprocess.run(["soffice", "--headless", "--calc", "--convert-to", "xlsx:Calc MS Excel 2007 XML", "--outdir", out] + paths[i:i + 12],
                       check=True, capture_output=True, timeout=600)
    return {p: os.path.join(out, os.path.basename(p)) for p in paths}


def read(path):
    S = openpyxl.load_workbook(path, data_only=True)["Resumo"]
    rows = {S.cell(r, 1).value: r for r in range(1, S.max_row + 1)}
    res = {}
    for key, label in OUTPUTS:
        r = rows[label]
        res[key] = {"A": S.cell(r, 2).value, "B": S.cell(r, 3).value}
    return res


# ------------------------------------------------------------------ casos
cases = {}
for c, nome in ((1, "Conservador"), (2, "Base"), (3, "Otimista")):
    cases[f"cen_{c}"] = {"cenario": c}
# estresses do Modelo B: e se as apostas centrais falharem?
STRESS = [
    ("B com churn igual ao do Modelo A", {"B_P_churn": 0.055, "B_G_churn": 0.055, "B_K_churn": 0.035}),
    ("B com metade das contas por especialista", {"B_P_esp": 35, "B_G_esp": 20, "B_K_esp": 8}),
    ("B sem receita por consulta", {"B_G_consultas": 0, "B_K_consultas": 0}),
    ("B com preço 30% menor (R$ 690 / R$ 1.250)", {"B_P_preco": 690, "B_G_preco": 1250, "B_K_preco": 3150}),
    ("B com tudo isso junto", {"B_P_churn": 0.055, "B_G_churn": 0.055, "B_K_churn": 0.035, "B_P_esp": 35, "B_G_esp": 20, "B_K_esp": 8,
                               "B_G_consultas": 0, "B_K_consultas": 0, "B_P_preco": 690, "B_G_preco": 1250, "B_K_preco": 3150}),
    ("Sem expansão para dentistas", {"exp_mes": 999}),
]
for i, (label, ov) in enumerate(STRESS):
    cases[f"st_{i}"] = dict(ov, cenario=2)

# sensibilidade (tornado) no cenário base: valor baixo e alto de cada premissa
TORNADO = {
    "A": [
        ("Churn do Consultório", "A_C_churn", 0.07, 0.04),
        ("Churn do Essencial", "A_E_churn", 0.09, 0.055),
        ("CAC na mídia paga", "A_cac", 6000, 3000),
        ("Preço do Consultório", "A_C_preco", 497, 897),
        ("Indicação (% da base por mês)", "indic", 0.005, 0.015),
        ("Parcerias (multiplicador)", "parc_mult", 0.6, 1.3),
        ("Custo direto do Consultório", "A_C_custo", 80, 30),
    ],
    "B": [
        ("Churn da Agenda Cheia", "B_G_churn", 0.05, 0.02),
        ("Churn da Presença", "B_P_churn", 0.055, 0.025),
        ("CAC na mídia paga", "B_cac", 7500, 4000),
        ("Preço da Agenda Cheia", "B_G_preco", 1290, 2290),
        ("Contas por especialista (Agenda Cheia)", "B_G_esp", 20, 60),
        ("Consultas atribuídas por cliente", "B_G_consultas", 0, 20),
        ("Custo dos agentes por conta (Agenda Cheia)", "B_G_custo", 400, 150),
        ("Indicação (% da base por mês)", "indic", 0.005, 0.015),
    ],
}
for m, items in TORNADO.items():
    for j, (label, key, lo, hi) in enumerate(items):
        cases[f"t{m}_{j}_lo"] = {"cenario": 2, key: lo}
        cases[f"t{m}_{j}_hi"] = {"cenario": 2, key: hi}

paths = {name: build(name, ov) for name, ov in cases.items()}
calc = recalc_all(list(paths.values()))
R = {name: read(calc[p]) for name, p in paths.items()}
json.dump(R, open(os.path.join(HERE, "resultados_sensibilidade.json"), "w"), indent=1, ensure_ascii=False)

# ------------------------------------------------------------------ grava no xlsx final
subprocess.run(["python3", BUILD, FINAL], check=True, capture_output=True)
wb = openpyxl.load_workbook(FINAL)
BOLD = Font(name="Arial", bold=True)
TITLE = Font(name="Arial", bold=True, size=14, color="1F2440")
H2 = Font(name="Arial", bold=True, color="FFFFFF")
HFILL = PatternFill("solid", fgColor="1F2440")
FMT = {"arr5": 'R$ #,##0', "arr7": 'R$ #,##0', "cap": 'R$ #,##0', "rpp": 'R$ #,##0', "mb5": '0.0%', "pen50": '0.0%', "ltvcac": '0.0"x"', "payback": '0.0', "eq": '0', "m50": '0', "cli5": '#,##0'}


def hdr(ws, r, labels):
    for i, lab in enumerate(labels):
        c = ws.cell(row=r, column=1 + i, value=lab)
        c.font = H2
        c.fill = HFILL
        c.alignment = Alignment(wrap_text=True, horizontal="left" if i == 0 else "center")


C = wb.create_sheet("Cenarios")
C.column_dimensions["A"].width = 52
for col in range(2, 16):
    C.column_dimensions[openpyxl.utils.get_column_letter(col)].width = 15
C.cell(1, 1, "Cenários e testes de estresse (valores gerados por sensibilidade_v2.py)").font = TITLE
r = 3
for m in ("A", "B"):
    C.cell(r, 1, "Modelo A: software por assinatura" if m == "A" else "Modelo B: agência operada por agentes").font = BOLD
    r += 1
    hdr(C, r, ["Indicador", "Conservador", "Base", "Otimista"]); r += 1
    for key, label in OUTPUTS:
        C.cell(r, 1, label)
        for i, c in enumerate((1, 2, 3)):
            cell = C.cell(r, 2 + i, R[f"cen_{c}"][key][m])
            cell.number_format = FMT[key]
        r += 1
    r += 1
C.cell(r, 1, "Testes de estresse do Modelo B (cenário base, uma aposta por vez)").font = BOLD
r += 1
hdr(C, r, ["Teste"] + [label for _, label in OUTPUTS]); r += 1
C.row_dimensions[r - 1].height = 60
for i, (label, _) in enumerate(STRESS):
    C.cell(r, 1, label)
    for j, (key, _) in enumerate(OUTPUTS):
        cell = C.cell(r, 2 + j, R[f"st_{i}"][key]["B"])
        cell.number_format = FMT[key]
    r += 1

T = wb.create_sheet("Sensibilidade")
T.column_dimensions["A"].width = 44
for col in range(2, 10):
    T.column_dimensions[openpyxl.utils.get_column_letter(col)].width = 17
T.cell(1, 1, "Sensibilidade no cenário base: quanto cada premissa mexe no resultado (valores gerados por script)").font = TITLE
r = 3
base = R["cen_2"]
for m in ("A", "B"):
    T.cell(r, 1, "Modelo A" if m == "A" else "Modelo B").font = BOLD
    r += 1
    hdr(T, r, ["Premissa (ordenada pelo impacto no capital e na receita)", "Valor baixo", "Valor alto", "ARR ano 7 (baixo)", "ARR ano 7 (alto)", "Diferença no ARR", "Capital necessário (baixo)", "Capital necessário (alto)", "LTV/CAC (baixo / alto)"]); r += 1
    rows = []
    for j, (label, key, lo, hi) in enumerate(TORNADO[m]):
        a, b = R[f"t{m}_{j}_lo"], R[f"t{m}_{j}_hi"]
        rows.append((abs((b["cap"][m] or 0) - (a["cap"][m] or 0)) + abs((b["arr7"][m] or 0) - (a["arr7"][m] or 0)) / 10, label, lo, hi, a, b))
    for amp, label, lo, hi, a, b in sorted(rows, reverse=True):
        T.cell(r, 1, label)
        T.cell(r, 2, lo)
        T.cell(r, 3, hi)
        for col, v, f in ((4, a["arr7"][m], FMT["arr7"]), (5, b["arr7"][m], FMT["arr7"]), (6, abs((b["arr7"][m] or 0) - (a["arr7"][m] or 0)), FMT["arr7"]), (7, a["cap"][m], FMT["cap"]), (8, b["cap"][m], FMT["cap"])):
            T.cell(r, col, v).number_format = f
        T.cell(r, 9, f"{a['ltvcac'][m]:.1f}x / {b['ltvcac'][m]:.1f}x")
        r += 1
    T.cell(r, 1, f"Base: ARR ano 7 = R$ {base['arr7'][m]:,.0f}; capital necessário = R$ {base['cap'][m]:,.0f}".replace(",", "."))
    r += 2

wb.save(FINAL)
print(json.dumps({k: v for k, v in R.items() if k.startswith(("cen_", "st_"))}, ensure_ascii=False, indent=0)[:6000])
for m in ("A", "B"):
    print("TORNADO", m)
    for j, (label, key, lo, hi) in enumerate(TORNADO[m]):
        a, b = R[f"t{m}_{j}_lo"], R[f"t{m}_{j}_hi"]
        print(f"  {label}: arr7 {a['arr7'][m]:,.0f} .. {b['arr7'][m]:,.0f} | cap {a['cap'][m]:,.0f} .. {b['cap'][m]:,.0f} | ltv/cac {a['ltvcac'][m]:.1f}..{b['ltvcac'][m]:.1f}")
