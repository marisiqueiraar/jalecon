"""Modelo financeiro do pivô para clínicas (Modelo C): agência operada por agentes, vendida a clínicas.

Mesma mecânica e mesmos critérios de corte do modelo v2.1 (build_model_v2.py), com o que muda na venda
para clínicas: planos por porte, venda ativa (duplas de SDR e executivo), receita de expansão
(médicos e unidades adicionais) e mercado endereçável contado em clínicas.
Rodar: python 10_v1/financeiro/build_model_clinicas.py [saida.xlsx] [overrides.json]
"""
import json
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter as L

OUT = sys.argv[1] if len(sys.argv) > 1 else "/home/user/jalecon/10_v1/financeiro/Jalecon_Modelo_Clinicas.xlsx"
OVR = json.load(open(sys.argv[2])) if len(sys.argv) > 2 else {}

N = 84  # meses (7 anos)
YEARS = 7
FONT = "Arial"
BLUE = Font(name=FONT, color="0000FF")
BLACK = Font(name=FONT, color="000000")
BOLD = Font(name=FONT, bold=True)
TITLE = Font(name=FONT, bold=True, size=14, color="1F2440")
H2 = Font(name=FONT, bold=True, color="FFFFFF")
HFILL = PatternFill("solid", fgColor="1F2440")
SFILL = PatternFill("solid", fgColor="E9EBF2")
KFILL = PatternFill("solid", fgColor="FBE9E0")
BRL = 'R$ #,##0;(R$ #,##0);-'
NUM = '#,##0;(#,##0);-'
NUM1 = '#,##0.0;(#,##0.0);-'
PCT = '0.0%;(0.0%);-'
MULT = '0.0"x"'

wb = Workbook()
A = {}  # chave -> referência absoluta da célula "usada"


def put(ws, r, c, v, font=BLACK, fmt=None, fill=None, bold=False):
    cell = ws.cell(row=r, column=c, value=v)
    cell.font = Font(name=FONT, bold=bold, color=font.color) if bold else font
    if fmt:
        cell.number_format = fmt
    if fill:
        cell.fill = fill
    return cell


def head(ws, r, labels, c0=1):
    for i, lab in enumerate(labels):
        cell = put(ws, r, c0 + i, lab, H2, fill=HFILL)
        cell.alignment = Alignment(horizontal="left" if i == 0 else "center", wrap_text=True)


def section(ws, r, text, width):
    for c in range(1, width + 1):
        ws.cell(row=r, column=c).fill = SFILL
    put(ws, r, 1, text, BOLD)


# ============================================================ Premissas
P = wb.active
P.title = "Premissas"
P.column_dimensions["A"].width = 58
P.column_dimensions["B"].width = 12
for col in "CDEF":
    P.column_dimensions[col].width = 14
P.column_dimensions["G"].width = 90
put(P, 1, 1, "Jalecon Clínicas (Modelo C) · Premissas", TITLE)
put(P, 2, 1, "Texto azul = premissa editável. A coluna F (Usado) segue o cenário escolhido em C3. A coluna G diz de onde veio cada número.")
put(P, 3, 1, "Cenário (1 = Conservador, 2 = Base, 3 = Otimista)", BOLD)
put(P, 3, 3, OVR.get("cenario", 2), BLUE, fill=KFILL)
put(P, 3, 4, '=CHOOSE(C3,"Conservador","Base","Otimista")', BOLD)
A["cenario_nome"] = "Premissas!$D$3"
row = 5


def scen(key, label, unit, cons, base, opt, fmt, src):
    """Premissa com três cenários; F = valor usado."""
    global row
    vals = OVR.get(key)
    if vals is not None:
        cons = base = opt = vals
    # multiplicadores usados pelo script de sensibilidade (mexem nos três cenários ao mesmo tempo)
    for suffix, mkey in (("_esp", "mult_esp"), ("_churn", "mult_churn"), ("_preco", "mult_preco"), ("cac_midia", "mult_cac")):
        if key.endswith(suffix) and key != "custo_esp" and mkey in OVR:
            cons, base, opt = (v * OVR[mkey] for v in (cons, base, opt))
    if key == "fech_dupla" and "mult_cac" in OVR:
        cons, base, opt = (v / OVR["mult_cac"] for v in (cons, base, opt))
    put(P, row, 1, label)
    put(P, row, 2, unit)
    for i, v in enumerate((cons, base, opt)):
        put(P, row, 3 + i, v, BLUE, fmt)
    put(P, row, 6, f"=CHOOSE($C$3,C{row},D{row},E{row})", BLACK, fmt)
    put(P, row, 7, src)
    A[key] = f"Premissas!$F${row}"
    row += 1


def yearly(key, label, unit, vals, fmt, src):
    """Decisão por ano (anos 1 a 7), igual nos três cenários. Colunas C..I."""
    global row
    vals = OVR.get(key, vals)
    put(P, row, 1, label)
    put(P, row, 2, unit)
    for i, v in enumerate(vals):
        put(P, row, 3 + i, v, BLUE, fmt)
    put(P, row, 10, src)
    A[key] = (row, 3)
    row += 1


def sec(text):
    global row
    row += 1
    section(P, row, text, 7)
    row += 1


def yhead():
    global row
    head(P, row, ["", ""] + [f"Ano {y}" for y in range(1, YEARS + 1)] + ["Fonte"])
    row += 1


def shead():
    global row
    head(P, row, ["Premissa", "Unidade", "Conservador", "Base", "Otimista", "Usado", "Fonte ou justificativa"])
    row += 1


sec("1. Gerais")
shead()
scen("lancamento", "Mês de lançamento comercial (antes: piloto concierge)", "mês", 5, 4, 4, NUM, "Como no v2.1")
scen("imp_baixo", "Impostos sobre a receita até R$ 4,8 mi nos últimos 12 meses", "%", 0.163, 0.163, 0.163, PCT, "Lucro presumido desde o início, como no v2.1")
scen("imp_alto", "Impostos sobre a receita acima de R$ 4,8 mi nos últimos 12 meses", "%", 0.163, 0.163, 0.163, PCT, "Lucro presumido para serviços (aprox.)")
scen("meios_pg", "Meios de pagamento e inadimplência", "% da receita", 0.04, 0.03, 0.025, PCT, "Boleto e Pix recorrentes; inclui inadimplência de PME")
scen("sam_cli", "Mercado endereçável: clínicas médicas privadas de 2 a 15 médicos que dependem de pacientes particulares", "clínicas", 10000, 20000, 35000, NUM, "pesquisa/clinicas.md: 30 a 70 mil clínicas privadas de 2 a 15 médicos (extrapolação do CNES da Bahia), das quais 30% a 50% investem em captação e têm particular relevante. Estimativa, sem fonte direta")
scen("exp_mes", "Expansão para clínicas odontológicas: mês de início", "mês", 37, 25, 25, NUM, "Exige base de regras do CFO 196/2019 e produto adaptado")
scen("exp_fator", "Expansão: aquisição adicional sobre a de clínicas médicas", "%", 0.3, 0.5, 0.8, PCT, "Clinicorp e Simples Dental mostram mercado grande e já disputado")
scen("sam_odonto", "Mercado endereçável: clínicas odontológicas com 2 ou mais dentistas", "clínicas", 8000, 15000, 25000, NUM, "pesquisa/clinicas.md: 191,7 mil CNPJs odontológicos e Clinicorp com mais de 30 mil clínicas; endereçáveis ao ticket de R$ 2.900 ou mais. Estimativa")
scen("exp_custo", "Custo único da expansão (regras, produto, jurídico)", "R$", 400000, 300000, 200000, BRL, "Como no v2.1")
scen("indic", "Indicação: novas clínicas por mês, como % da base ativa", "%", 0.004, 0.008, 0.012, PCT, "Gestores de clínica conversam menos entre si do que médicos; abaixo do v2.1 (1%)")
scen("parc_mult", "Parcerias: multiplicador sobre o plano de parcerias", "x", 0.6, 1.0, 1.3, MULT, "Softwares de gestão (Feegow, iClinic, Clinicorp), contadores e consultorias de clínicas")

sec("2. Decisões por ano (iguais nos três cenários)")
yhead()
yearly("midia", "Orçamento de mídia paga por mês (gera reuniões com gestores)", "R$", [15000, 30000, 50000, 80000, 110000, 140000, 170000], BRL, "Decisão; metade do Modelo B, porque a venda é ativa")
yearly("duplas", "Duplas de venda ativa (SDR e executivo)", "duplas", [1, 2, 4, 6, 8, 10, 12], NUM, "Decisão; uma dupla nova produz 50% no ano em que entra")
yearly("parc", "Novas clínicas por parcerias por mês (antes do multiplicador)", "clínicas", [0, 1, 2, 3, 4, 5, 6], NUM, "Suposição sem evidência de canal; integração com softwares de gestão só a partir do ano 2")
yearly("hc_prod", "Produto, engenharia e IA (pessoas)", "pessoas", [4, 6, 8, 10, 12, 13, 14], NUM, "Menos que no v2.1: um só tipo de cliente")
yearly("hc_mkt", "Marketing, pré-venda inbound e sucesso do cliente (pessoas)", "pessoas", [1, 2, 3, 4, 5, 6, 7], NUM, "Fecha os contatos que vêm da mídia")
yearly("hc_ga", "Administrativo, jurídico e conformidade (pessoas)", "pessoas", [1, 2, 3, 4, 6, 7, 8], NUM, "Plano de contratação")
yearly("prolabore", "Pró-labore por fundador por mês", "R$", [5000, 10000, 15000, 15000, 18000, 18000, 20000], BRL, "Como no v2.1 (6 fundadores)")

sec("3. Custos de pessoas e despesas gerais")
shead()
scen("custo_prod", "Custo mensal por pessoa de produto e engenharia (com encargos)", "R$", 20000, 18000, 16000, BRL, "Como no v2.1")
scen("custo_dupla", "Custo mensal de uma dupla de venda (SDR e executivo, com comissão)", "R$", 24000, 22000, 20000, BRL, "SDR de R$ 7 mil e executivo de R$ 13 mil com encargos, mais comissão")
scen("custo_mkt", "Custo mensal por pessoa de marketing e sucesso do cliente", "R$", 11000, 10000, 9000, BRL, "Como vendas e marketing no v2.1")
scen("custo_ga", "Custo mensal por pessoa administrativa e jurídica", "R$", 13000, 12000, 11000, BRL, "Como no v2.1")
scen("fundadores", "Número de fundadores", "pessoas", 6, 6, 6, NUM, "Grupo 1")
scen("outras_fixo", "Outras despesas fixas por mês (ferramentas, jurídico, contabilidade)", "R$", 20000, 15000, 12000, BRL, "Como no v2.1")
scen("outras_pct", "Outras despesas variáveis", "% da receita", 0.04, 0.03, 0.025, PCT, "Como no v2.1")
scen("marca_pct", "Marca, eventos e conteúdo próprio", "% da mídia", 0.10, 0.08, 0.06, PCT, "Feiras de gestão de clínicas e conteúdo para diretores técnicos")

sec("4. Aquisição")
shead()
scen("cac_midia", "Custo de mídia por clínica fechada via inbound", "R$", 30000, 22000, 15000, BRL, "Calibrado para o CAC total do base ficar em cerca de 9,5 meses de contribuição (mesmo método da regra pré-registrada do v2.1); a confirmar no teste de aquisição")
scen("fech_dupla", "Clínicas fechadas por dupla de venda ativa por mês (em ritmo pleno)", "clínicas", 0.7, 1.0, 1.5, NUM1, "Calibrado como o CAC de mídia: R$ 22 mil por fechamento no base. Equivale a cerca de R$ 640 mil de receita anual nova por executivo")
scen("comissao", "Custo por clínica vinda de parceria (comissão)", "R$", 11000, 9000, 7000, BRL, "Cerca de 20% da receita do primeiro ano, como em programas de revenda de software")
scen("custo_ind", "Custo por clínica indicada (mês grátis)", "R$", 4300, 4300, 4300, BRL, "Um mês grátis na mensalidade média ponderada")

PLANS = [
    # chave, nome, porte, preço, setup, mix (c,b,o), churn (c,b,o), custo direto (c,b,o), contas por especialista (c,b,o), mês de início
    ("E", "Clínica Essencial", "2 a 4 médicos, 1 unidade", 2900, 2900, (0.55, 0.50, 0.45), (0.045, 0.03, 0.02), (450, 350, 280), (16, 25, 35), 1),
    ("P", "Clínica Pro", "5 a 10 médicos", 4900, 4900, (0.38, 0.40, 0.42), (0.035, 0.025, 0.016), (700, 550, 450), (10, 16, 24), 1),
    ("R", "Rede", "11 ou mais médicos ou 2 ou mais unidades", 8900, 7900, (0.07, 0.10, 0.13), (0.03, 0.02, 0.012), (1200, 950, 800), (6, 9, 14), 13),
]
sec("5. Planos")
shead()
scen("custo_esp", "Custo mensal do especialista que cuida das contas (gerente de conta sênior)", "R$", 13000, 12000, 11000, BRL, "Mais sênior que no Modelo B (R$ 10 mil): o interlocutor é o gestor da clínica")
scen("impl", "Custo de implantação por clínica nova", "R$", 2500, 1800, 1200, BRL, "Marca, site, perfis no Google de cada médico, integração do WhatsApp e da agenda: 12 a 20 horas de especialista mais ferramentas")
scen("exp_taxa", "Receita de expansão: aumento da mensalidade da base por mês (médicos e unidades novas)", "% da mensalidade", 0.003, 0.006, 0.01, PCT, "Suposição; clínica que contrata médico novo paga o perfil adicional")
scen("exp_cogs", "Custo do serviço da receita de expansão", "% dessa receita", 0.45, 0.35, 0.25, PCT, "Médico adicional exige perfil, conteúdo e tempo de especialista")
for k, name, porte, price, setup, mix, churn, custo, esp, start in PLANS:
    scen(f"{k}_preco", f"{name} ({porte}): preço mensal", "R$", price, price, price, BRL, "Abaixo de agências para clínicas (R$ 5 mil a R$ 15 mil sem mídia) e acima da Medleads (R$ 3.997)")
    scen(f"{k}_setup", f"{name}: taxa de implantação", "R$", setup, setup, setup, BRL, "Uma mensalidade, como agências")
    scen(f"{k}_mix", f"{name}: participação nas clínicas novas", "%", *mix, PCT, "Suposição" + (" (a partir do mês 13; antes vai para a Clínica Pro)" if start > 1 else ""))
    scen(f"{k}_churn", f"{name}: churn mensal", "%", *churn, PCT, "Pesquisa de clínicas, seção 4: software vertical e agências para clínicas; contrato mensal")
    scen(f"{k}_custo", f"{name}: custo direto por clínica por mês (agentes, WhatsApp, ferramentas)", "R$", *custo, BRL, "IA de cerca de R$ 20 por médico (v2.1) mais API do WhatsApp, site, design e agenda")
    scen(f"{k}_esp", f"{name}: clínicas por especialista humano", "clínicas", *esp, NUM, "Com 120 horas produtivas por mês: base de cerca de 5, 7,5 e 13 horas por clínica")

sec("6. Critérios de corte do plano v1 (fixados em 19/10/2026; os mesmos para o pivô)")
shead()
scen("crit_mb", "H6: margem bruta mínima", "%", 0.70, 0.70, 0.70, PCT, "Plano v1, seção 2")
scen("crit_ltvcac", "H6: LTV/CAC mínimo", "x", 3, 3, 3, MULT, "Plano v1, seção 2. Aqui medido com o CAC total de vendas e marketing (mais exigente que o CAC de mídia do v2.1)")
scen("crit_payback", "H6: recuperação máxima do CAC", "meses", 12, 12, 12, NUM, "Plano v1, seção 2")
scen("crit_arr", "H7: receita recorrente anual mínima em até 7 anos", "R$", 50000000, 50000000, 50000000, BRL, "Plano v1, seção 2")
scen("crit_pen", "H7: participação máxima do mercado endereçável quando atingir esse valor", "%", 0.03, 0.03, 0.03, PCT, "Plano v1, seção 2")


def ycell(key, year_ref):
    r, c0 = A[key]
    return f"INDEX(Premissas!$C${r}:$I${r},1,{year_ref})"


# ============================================================ Aba mensal
ws = wb.create_sheet("Modelo C")
ws.column_dimensions["A"].width = 52
ws.column_dimensions["B"].width = 4
for c in range(3, 3 + N):
    ws.column_dimensions[L(c)].width = 11
ws.freeze_panes = "C3"
put(ws, 1, 1, "Modelo C: agência operada por agentes para clínicas · mês a mês", TITLE)
R = {}
r = 2
pending = []


def line(key, label, fn, fmt=NUM, bold=False, fill=None):
    global r
    put(ws, r, 1, label, BOLD if bold else BLACK, fill=fill)
    pending.append((r, fn, fmt, bold, fill))
    R[key] = r
    r += 1


def ref(key, col):
    return f"{L(col)}{R[key]}"


def prev(key, col):
    return f"{L(col - 1)}{R[key]}" if col > 3 else "0"


def gap(title):
    global r
    r += 1
    section(ws, r, title, 2)
    r += 1


line("mes", "Mês", lambda i, c: i + 1, NUM, bold=True)
line("ano", "Ano", lambda i, c: f"=ROUNDUP({ref('mes', c)}/12,0)")
line("ativo", "Venda ativa (1 a partir do lançamento)", lambda i, c: f"=IF({ref('mes', c)}>={A['lancamento']},1,0)")
line("exp", "Fator de expansão para clínicas odontológicas", lambda i, c: f"=IF({ref('mes', c)}>={A['exp_mes']},1+{A['exp_fator']},1)", MULT)
gap("Aquisição")
line("midia", "Mídia paga no mês", lambda i, c: f"={ycell('midia', ref('ano', c))}*{ref('ativo', c)}*{ref('exp', c)}", BRL)
line("n_midia", "Clínicas novas por mídia (inbound)", lambda i, c: f"={ref('midia', c)}/{A['cac_midia']}", NUM1)
line("duplas", "Duplas de venda contratadas", lambda i, c: f"={ycell('duplas', ref('ano', c))}*{ref('exp', c)}", NUM1)
line("duplas_ef", "Duplas em ritmo pleno (as novas do ano produzem 50%)",
     lambda i, c: f"=({ycell('duplas', ref('ano', c))}+IF({ref('ano', c)}>1,{ycell('duplas', ref('ano', c) + '-1')},0))/2*{ref('exp', c)}", NUM1)
line("n_ativa", "Clínicas novas por venda ativa", lambda i, c: f"={ref('duplas_ef', c)}*{A['fech_dupla']}*{ref('ativo', c)}", NUM1)
line("n_parc", "Clínicas novas por parcerias", lambda i, c: f"={ycell('parc', ref('ano', c))}*{A['parc_mult']}*{ref('ativo', c)}*{ref('exp', c)}", NUM1)
line("n_ind", "Clínicas novas por indicação", lambda i, c: f"={prev('tot_ativos', c)}*{A['indic']}*{ref('ativo', c)}" if c > 3 else "=0", NUM1)
line("n_tot", "Clínicas novas no mês", lambda i, c: f"={ref('n_midia', c)}+{ref('n_ativa', c)}+{ref('n_parc', c)}+{ref('n_ind', c)}", NUM1, bold=True)

for k, name, porte, *_rest in PLANS:
    start = _rest[-1]
    gap(f"Plano {name}")
    mixref = A[f"{k}_mix"]
    if k == "R":
        line(f"mix_{k}", "Participação nas novas", lambda i, c, mixref=mixref: f"=IF({ref('mes', c)}>=13,{mixref},0)", PCT)
    elif k == "P":
        line(f"mix_{k}", "Participação nas novas", lambda i, c, mixref=mixref: f"={mixref}+IF({ref('mes', c)}<13,{A['R_mix']},0)", PCT)
    else:
        line(f"mix_{k}", "Participação nas novas", lambda i, c, mixref=mixref: f"={mixref}", PCT)
    line(f"novos_{k}", "Novas", lambda i, c, k=k: f"={ref('n_tot', c)}*{ref('mix_' + k, c)}", NUM1)
    line(f"perd_{k}", "Canceladas", lambda i, c, k=k: f"={prev('ativos_' + k, c)}*{A[f'{k}_churn']}", NUM1)
    line(f"ativos_{k}", "Ativas no fim do mês", lambda i, c, k=k: f"={prev('ativos_' + k, c)}+{ref('novos_' + k, c)}-{ref('perd_' + k, c)}", NUM1, bold=True)

gap("Clientes e receita")
line("tot_ativos", "Clínicas ativas", lambda i, c: "=" + "+".join(ref("ativos_" + p[0], c) for p in PLANS), NUM, bold=True)
line("tot_perd", "Clínicas canceladas no mês", lambda i, c: "=" + "+".join(ref("perd_" + p[0], c) for p in PLANS), NUM1)
line("sam", "Mercado endereçável no mês", lambda i, c: f"={A['sam_cli']}+IF({ref('mes', c)}>={A['exp_mes']},{A['sam_odonto']},0)", NUM)
line("pen", "Participação no mercado endereçável", lambda i, c: f"={ref('tot_ativos', c)}/{ref('sam', c)}", PCT)
line("rec_fixa", "Receita de mensalidades dos planos", lambda i, c: "=" + "+".join(f"{ref('ativos_' + p[0], c)}*{A[p[0] + '_preco']}" for p in PLANS), BRL)
line("rec_exp", "Receita de expansão (médicos e unidades adicionais)",
     lambda i, c: (f"={prev('rec_exp', c)}*(1-IF({prev('tot_ativos', c)}>0,{ref('tot_perd', c)}/{prev('tot_ativos', c)},0))"
                   f"+{prev('rec_fixa', c)}*{A['exp_taxa']}") if c > 3 else "=0", BRL)
line("rec_setup", "Receita de implantação", lambda i, c: "=" + "+".join(f"{ref('novos_' + p[0], c)}*{A[p[0] + '_setup']}" for p in PLANS), BRL)
line("rec", "Receita bruta", lambda i, c: f"={ref('rec_fixa', c)}+{ref('rec_exp', c)}+{ref('rec_setup', c)}", BRL, bold=True)
line("mrr", "Receita recorrente mensal (mensalidades + expansão)", lambda i, c: f"={ref('rec_fixa', c)}+{ref('rec_exp', c)}", BRL)
line("arr", "Receita recorrente anualizada (MRR × 12)", lambda i, c: f"={ref('mrr', c)}*12", BRL, bold=True)
line("arpa", "Mensalidade média por clínica", lambda i, c: f"=IF({ref('tot_ativos', c)}>0,{ref('mrr', c)}/{ref('tot_ativos', c)},0)", BRL)
LAST = L(2 + N)


def trailing(c):
    # receita dos 12 meses anteriores; forma invariante para poder copiar a coluna para a direita
    return (f'SUMIFS($C${R["rec"]}:${LAST}${R["rec"]},$C${R["mes"]}:${LAST}${R["mes"]},">="&({ref("mes", c)}-12),'
            f'$C${R["mes"]}:${LAST}${R["mes"]},"<"&{ref("mes", c)})')


line("imp", "Impostos sobre a receita", lambda i, c: f"={ref('rec', c)}*IF({trailing(c)}>4800000,{A['imp_alto']},{A['imp_baixo']})", BRL)
line("mpg", "Meios de pagamento e inadimplência", lambda i, c: f"={ref('rec', c)}*{A['meios_pg']}", BRL)
line("rec_liq", "Receita líquida", lambda i, c: f"={ref('rec', c)}-{ref('imp', c)}-{ref('mpg', c)}", BRL, bold=True)

gap("Custo do serviço")
line("cogs_dir", "Custo direto (agentes, WhatsApp, ferramentas, implantação)",
     lambda i, c: "=" + "+".join(f"{ref('ativos_' + p[0], c)}*{A[p[0] + '_custo']}" for p in PLANS) + f"+{ref('n_tot', c)}*{A['impl']}", BRL)
line("esp_n", "Especialistas humanos nas contas (pessoas)", lambda i, c: "=" + "+".join(f"{ref('ativos_' + p[0], c)}/{A[p[0] + '_esp']}" for p in PLANS), NUM1)
line("cogs_esp", "Custo dos especialistas", lambda i, c: f"={ref('esp_n', c)}*{A['custo_esp']}", BRL)
line("cogs_exp", "Custo do serviço da receita de expansão", lambda i, c: f"={ref('rec_exp', c)}*{A['exp_cogs']}", BRL)
line("cogs", "Custo do serviço total", lambda i, c: f"={ref('cogs_dir', c)}+{ref('cogs_esp', c)}+{ref('cogs_exp', c)}", BRL, bold=True)
line("lb", "Lucro bruto", lambda i, c: f"={ref('rec_liq', c)}-{ref('cogs', c)}", BRL, bold=True)
line("mb", "Margem bruta (sobre a receita líquida)", lambda i, c: f"=IF({ref('rec_liq', c)}>0,{ref('lb', c)}/{ref('rec_liq', c)},0)", PCT)

gap("Despesas operacionais")
line("d_prod", "Produto, engenharia e IA", lambda i, c: f"={ycell('hc_prod', ref('ano', c))}*{A['custo_prod']}", BRL)
line("d_duplas", "Duplas de venda ativa", lambda i, c: f"={ref('duplas', c)}*{A['custo_dupla']}", BRL)
line("d_mkt", "Marketing, pré-venda inbound e sucesso do cliente", lambda i, c: f"={ycell('hc_mkt', ref('ano', c))}*{A['custo_mkt']}", BRL)
line("d_ga", "Administrativo, jurídico e conformidade", lambda i, c: f"={ycell('hc_ga', ref('ano', c))}*{A['custo_ga']}", BRL)
line("d_fund", "Pró-labore dos fundadores", lambda i, c: f"={ycell('prolabore', ref('ano', c))}*{A['fundadores']}", BRL)
line("d_midia", "Mídia paga", lambda i, c: f"={ref('midia', c)}", BRL)
line("d_parc", "Comissões de parceria e meses grátis de indicação", lambda i, c: f"={ref('n_parc', c)}*{A['comissao']}+{ref('n_ind', c)}*{A['custo_ind']}", BRL)
line("d_marca", "Marca e eventos", lambda i, c: f"={ref('midia', c)}*{A['marca_pct']}", BRL)
line("d_outras", "Outras despesas", lambda i, c: f"={A['outras_fixo']}+{ref('rec', c)}*{A['outras_pct']}", BRL)
line("d_exp", "Custo único da expansão", lambda i, c: f"=IF({ref('mes', c)}={A['exp_mes']},{A['exp_custo']},0)", BRL)
line("opex", "Despesas operacionais totais",
     lambda i, c: "=" + "+".join(ref(k, c) for k in ("d_prod", "d_duplas", "d_mkt", "d_ga", "d_fund", "d_midia", "d_parc", "d_marca", "d_outras", "d_exp")), BRL, bold=True)
line("sm", "Gasto total de aquisição (mídia, duplas, marketing, comissões, marca)",
     lambda i, c: "=" + "+".join(ref(k, c) for k in ("d_midia", "d_duplas", "d_mkt", "d_parc", "d_marca")), BRL)

gap("Resultado e caixa")
line("ebitda", "EBITDA", lambda i, c: f"={ref('lb', c)}-{ref('opex', c)}", BRL, bold=True, fill=KFILL)
line("caixa", "Caixa acumulado sem captação", lambda i, c: f"={prev('caixa', c)}+{ref('ebitda', c)}", BRL, bold=True)
line("eq_flag", "Mês de equilíbrio candidato (EBITDA não volta a ficar negativo)",
     lambda i, c: f'=IF(COUNTIF({L(c)}{R["ebitda"]}:${LAST}${R["ebitda"]},"<0")=0,{ref("mes", c)},"")', NUM)
line("arr_flag", "Mês em que a receita anualizada passa do critério H7",
     lambda i, c: f'=IF(AND({ref("arr", c)}>={A["crit_arr"]},{prev("arr", c)}<{A["crit_arr"]}),{ref("mes", c)},"")', NUM)
line("hc", "Pessoas no total (fundadores, equipe e especialistas)",
     lambda i, c: (f"={A['fundadores']}+{ycell('hc_prod', ref('ano', c))}+2*{ref('duplas', c)}+{ycell('hc_mkt', ref('ano', c))}"
                   f"+{ycell('hc_ga', ref('ano', c))}+{ref('esp_n', c)}"), NUM1)
for rr, fn, fmt, bold, fill in pending:
    for i in range(N):
        put(ws, rr, 3 + i, fn(i, 3 + i), BLACK, fmt, fill=fill, bold=bold)


def rng(key):
    return f"'Modelo C'!$C${R[key]}:${LAST}${R[key]}"


def at(key, month_expr):
    return f"INDEX({rng(key)},1,{month_expr})"


def ysum(key, y):
    return f"SUMIF('Modelo C'!$C${R['ano']}:${LAST}${R['ano']},{y},{rng(key)})"


# ============================================================ Unidade
U = wb.create_sheet("Unidade", 1)
U.column_dimensions["A"].width = 44
for c in range(2, 11):
    U.column_dimensions[L(c)].width = 15
put(U, 1, 1, "Economia por clínica, por plano (cenário selecionado)", TITLE)
put(U, 2, 1, "Contribuição = mensalidade menos impostos (alíquota cheia), meios de pagamento, custo direto e parte do especialista. LTV = contribuição / churn + taxa de implantação - custo de implantação (sem receita de expansão).")
head(U, 4, ["Plano", "Mensalidade", "Contribuição mensal", "Margem de contribuição", "Churn mensal", "Vida média (meses)", "LTV", "LTV/CAC (CAC total do ano 5)", "Recuperação do CAC (meses)", "Margem bruta do plano"])
UREF = {}
r = 5
for k, name, *_ in PLANS:
    put(U, r, 1, name)
    put(U, r, 2, f"={A[k + '_preco']}", fmt=BRL)
    put(U, r, 3, f"=B{r}*(1-{A['imp_alto']}-{A['meios_pg']})-{A[k + '_custo']}-{A['custo_esp']}/{A[k + '_esp']}", fmt=BRL)
    put(U, r, 4, f"=C{r}/B{r}", fmt=PCT)
    put(U, r, 5, f"={A[k + '_churn']}", fmt=PCT)
    put(U, r, 6, f"=1/E{r}", fmt=NUM1)
    put(U, r, 7, f"=C{r}/E{r}+{A[k + '_setup']}-{A['impl']}", fmt=BRL)
    put(U, r, 8, f"=G{r}/$B$11", fmt=MULT)
    put(U, r, 9, f"=$B$11/C{r}", fmt=NUM1)
    put(U, r, 10, f"=C{r}/(B{r}*(1-{A['imp_alto']}-{A['meios_pg']}))", fmt=PCT)
    UREF[k] = r
    r += 1
put(U, r, 1, "Ponderado pelo mix de clínicas novas", BOLD)
for col in (2, 3, 7):
    put(U, r, col, "=" + "+".join(f"{A[p[0] + '_mix']}*{L(col)}{UREF[p[0]]}" for p in PLANS), fmt=BRL, bold=True)
put(U, r, 8, f"=G{r}/$B$11", fmt=MULT, bold=True)
put(U, r, 9, f"=$B$11/C{r}", fmt=NUM1, bold=True)
UREF["pond"] = r
put(U, 11, 1, "CAC total no ano 5 (gasto de aquisição / clínicas novas)", BOLD)
put(U, 11, 2, f"={ysum('sm', 5)}/{ysum('n_tot', 5)}", fmt=BRL, bold=True)
put(U, 12, 1, "CAC só de mídia (comparável ao v2.1)")
put(U, 12, 2, f"={A['cac_midia']}", fmt=BRL)
put(U, 13, 1, "CAC total no ano 2")
put(U, 13, 2, f"={ysum('sm', 2)}/{ysum('n_tot', 2)}", fmt=BRL)

# ============================================================ Resumo
S = wb.create_sheet("Resumo", 0)
S.column_dimensions["A"].width = 60
for c in range(2, 10):
    S.column_dimensions[L(c)].width = 14
put(S, 1, 1, "Jalecon Clínicas (Modelo C) · Resumo", TITLE)
put(S, 2, 1, "Cenário:", BOLD)
put(S, 2, 2, f"={A['cenario_nome']}", BOLD)
put(S, 3, 1, "Troque o cenário na aba Premissas, célula C3. Valores em reais nominais.")
r = 5
section(S, r, "Ano a ano", 8); r += 1
head(S, r, ["Indicador"] + [f"Ano {y}" for y in range(1, YEARS + 1)]); r += 1
rows = [
    ("Clínicas ativas (fim do ano)", "tot_ativos", "end", NUM),
    ("Clínicas novas no ano", "n_tot", "sum", NUM),
    ("Receita bruta", "rec", "sum", BRL),
    ("Receita recorrente anualizada (fim do ano)", "arr", "end", BRL),
    ("Mensalidade média por clínica (fim do ano)", "arpa", "end", BRL),
    ("Lucro bruto", "lb", "sum", BRL),
    ("EBITDA", "ebitda", "sum", BRL),
    ("Caixa acumulado sem captação (fim do ano)", "caixa", "end", BRL),
    ("Pessoas (fim do ano)", "hc", "end", NUM),
    ("Participação no mercado endereçável (fim do ano)", "pen", "end", PCT),
]
for label, key, how, fmt in rows:
    put(S, r, 1, label)
    for y in range(1, YEARS + 1):
        put(S, r, 1 + y, f"={ysum(key, y)}" if how == "sum" else f"={at(key, 12 * y)}", fmt=fmt)
    if key == "lb":
        r += 1
        put(S, r, 1, "Margem bruta (sobre a receita líquida)")
        for y in range(1, YEARS + 1):
            put(S, r, 1 + y, f"=IFERROR({ysum('lb', y)}/{ysum('rec_liq', y)},0)", fmt=PCT)
    r += 1
r += 1
section(S, r, "Indicadores-chave e critérios de corte", 8); r += 1
head(S, r, ["Indicador", "Modelo C", "Critério", "Passa?"]); r += 1


def keyrow(label, f, fmt, crit=None, test=None):
    global r
    put(S, r, 1, label)
    put(S, r, 2, f, fmt=fmt)
    if crit is not None:
        put(S, r, 3, crit, fmt=fmt)
    if test:
        put(S, r, 4, test.format(v=f"B{r}"))
    r += 1
    return r - 1


keyrow("Receita recorrente anualizada no ano 5", f"={at('arr', 60)}", BRL)
keyrow("Receita recorrente anualizada no ano 7", f"={at('arr', 84)}", BRL)
keyrow("Mês de equilíbrio (EBITDA positivo daí em diante)", f'=IFERROR(1/(1/MIN({rng("eq_flag")})),"não chega")', NUM)
keyrow("Capital necessário (pior caixa acumulado)", f"=-MIN(0,MIN({rng('caixa')}))", BRL)
keyrow("Margem bruta no ano 5", f"={ysum('lb', 5)}/{ysum('rec_liq', 5)}", PCT, f"={A['crit_mb']}", '=IF({v}>=' + A["crit_mb"] + ',"Sim","Não")')
pr = UREF["pond"]
keyrow("LTV/CAC ponderado pelo mix, com CAC total do ano 5", f"=Unidade!H{pr}", MULT, f"={A['crit_ltvcac']}", '=IF({v}>=' + A["crit_ltvcac"] + ',"Sim","Não")')
keyrow("Recuperação do CAC total, ponderada (meses)", f"=Unidade!I{pr}", NUM1, f"={A['crit_payback']}", '=IF({v}<=' + A["crit_payback"] + ',"Sim","Não")')
keyrow("LTV/CAC ponderado com CAC só de mídia (comparável ao v2.1)", f"=Unidade!G{pr}/Unidade!B12", MULT)
keyrow("CAC total no ano 5", "=Unidade!B11", BRL)
m50 = keyrow("Mês em que a receita anualizada chega a R$ 50 mi", f'=IFERROR(1/(1/MIN({rng("arr_flag")})),"não chega")', NUM)
keyrow("Participação no mercado nesse mês", f'=IF(ISNUMBER(B{m50}),{at("pen", f"B{m50}")},"")', PCT, f"={A['crit_pen']}",
       '=IF(ISNUMBER({v}),IF({v}<=' + A["crit_pen"] + ',"Sim","Não"),"Não chega a R$ 50 mi")')
keyrow("Receita de expansão no ano 5, como % das mensalidades", f"={ysum('rec_exp', 5)}/{ysum('rec_fixa', 5)}", PCT)
keyrow("Receita por pessoa no ano 5", f"={ysum('rec', 5)}/{at('hc', 60)}", BRL)
keyrow("Clínicas ativas no ano 5", f"={at('tot_ativos', 60)}", NUM)
keyrow("Especialistas nas contas no ano 5", f"={at('esp_n', 60)}", NUM1)
r += 1
put(S, r, 1, "Como ler: H6 exige as três linhas com 'Sim' (margem, LTV/CAC e recuperação do CAC). H7 exige R$ 50 mi de receita anualizada em até 7 anos com no máximo 3% do mercado endereçável.")
r += 1
put(S, r, 1, "As abas Cenarios e Sensibilidade são geradas pelo script sensibilidade_clinicas.py: são valores, não fórmulas.")

# ============================================================ Leia-me
LM = wb.create_sheet("Leia-me", 1)
LM.column_dimensions["A"].width = 130
texts = [
    ("Jalecon Clínicas (Modelo C) · Como usar este modelo", TITLE),
    ("Pivô avaliado depois do memorando 08: a mesma agência operada por agentes de IA, vendida a clínicas privadas de 2 a 15 médicos (e redes), em vez de médicos individuais.", BLACK),
    ("O que muda em relação ao Modelo B do v2.1: planos por porte da clínica (R$ 2.900, R$ 4.900 e R$ 8.900), venda ativa com duplas de SDR e executivo, receita de expansão quando a clínica contrata médicos ou abre unidades, e mercado contado em clínicas.", BLACK),
    ("O que não muda: impostos, critérios de corte (H6 e H7), mecânica de caixa e cobrança por consulta desligada (art. 59 do Código de Ética Médica).", BLACK),
    ("Critério de LTV/CAC mais exigente que no v2.1: usa o CAC total (mídia, duplas, marketing, comissões e marca divididos pelas clínicas novas do ano 5), e não só o CAC de mídia. O LTV não inclui receita de expansão.", BLACK),
    ("Abas: Resumo, Leia-me, Unidade (economia por clínica), Premissas, Modelo C (84 meses), Cenarios e Sensibilidade (geradas por script).", BLACK),
    ("Limitações: tamanho do mercado, churn, produtividade de vendas e horas por clínica são referências de mercado, não medição própria. Não há sazonalidade nem reforma tributária.", BLACK),
]
for i, (t, f) in enumerate(texts, start=1):
    c = put(LM, i, 1, t, f)
    c.alignment = Alignment(wrap_text=True)

wb.save(OUT)
print("ok", OUT)
