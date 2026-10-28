"""Cenários, testes de estresse, sensibilidade e condições mínimas do Modelo C (clínicas).

Recalcula cada caso no LibreOffice e grava as abas Cenarios e Sensibilidade (valores) no xlsx final,
mais resultados_clinicas.json. Uso: python 10_v1/financeiro/sensibilidade_clinicas.py
"""
import json
import os
import subprocess
import tempfile

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

HERE = os.path.dirname(os.path.abspath(__file__))
BUILD = os.path.join(HERE, "build_model_clinicas.py")
FINAL = os.path.join(HERE, "Jalecon_Modelo_Clinicas.xlsx")
WORK = tempfile.mkdtemp(prefix="jalecon_c_")

OUTPUTS = [
    ("arr5", "Receita recorrente anualizada no ano 5"),
    ("arr7", "Receita recorrente anualizada no ano 7"),
    ("eq", "Mês de equilíbrio (EBITDA positivo daí em diante)"),
    ("cap", "Capital necessário (pior caixa acumulado)"),
    ("mb5", "Margem bruta no ano 5"),
    ("ltvcac", "LTV/CAC ponderado pelo mix, com CAC total do ano 5"),
    ("payback", "Recuperação do CAC total, ponderada (meses)"),
    ("cac5", "CAC total no ano 5"),
    ("m50", "Mês em que a receita anualizada chega a R$ 50 mi"),
    ("pen50", "Participação no mercado nesse mês"),
    ("cli5", "Clínicas ativas no ano 5"),
]


def build(name, overrides):
    xlsx = os.path.join(WORK, name + ".xlsx")
    ovr = os.path.join(WORK, name + ".json")
    json.dump(overrides, open(ovr, "w"))
    subprocess.run(["python3", BUILD, xlsx, ovr], check=True, capture_output=True)
    return xlsx


def recalc_all(paths):
    out = os.path.join(WORK, "out")
    os.makedirs(out, exist_ok=True)
    for i in range(0, len(paths), 12):
        subprocess.run(["soffice", "--headless", "--calc", "--convert-to", "xlsx:Calc MS Excel 2007 XML", "--outdir", out] + paths[i:i + 12],
                       check=True, capture_output=True, timeout=900)
    return {p: os.path.join(out, os.path.basename(p)) for p in paths}


def read(path):
    S = openpyxl.load_workbook(path, data_only=True)["Resumo"]
    rows = {S.cell(r, 1).value: r for r in range(1, S.max_row + 1)}
    return {key: S.cell(rows[label], 2).value for key, label in OUTPUTS}


def num(v):
    return v if isinstance(v, (int, float)) else None


def passes(res):
    h6 = (res["mb5"] or 0) >= 0.70 and (res["ltvcac"] or 0) >= 3 and (num(res["payback"]) or 99) <= 12
    h7 = num(res["m50"]) is not None and (num(res["pen50"]) or 1) <= 0.03
    return h6, h7


cases = {f"cen_{c}": {"cenario": c} for c in (1, 2, 3)}
STRESS = [
    ("Churn 50% maior em todos os planos", {"mult_churn": 1.5}),
    ("CAC 50% maior (mídia e venda ativa)", {"mult_cac": 1.5}),
    ("Metade das clínicas por especialista", {"mult_esp": 0.5}),
    ("O dobro de clínicas por especialista", {"mult_esp": 2.0}),
    ("Preços 20% menores", {"mult_preco": 0.8}),
    ("Sem receita de expansão", {"exp_taxa": 0}),
    ("Sem clínicas odontológicas", {"exp_mes": 999}),
    ("Mercado de clínicas médicas no conservador e sem odontologia", {"sam_cli": 10000, "exp_mes": 999}),
    ("CAC da pesquisa de clínicas (cerca de R$ 12 mil, 40% menor)", {"mult_cac": 0.6}),
    ("Churn e CAC 50% maiores juntos", {"mult_churn": 1.5, "mult_cac": 1.5}),
    ("Só clínicas de 5 ou mais médicos (sem o Essencial; mercado pela metade)", {"E_mix": 0, "P_mix": 0.8, "R_mix": 0.2, "sam_cli": 10000, "sam_odonto": 7500}),
    ("Só 5 ou mais médicos, com CAC 30% maior (venda mais longa)", {"E_mix": 0, "P_mix": 0.8, "R_mix": 0.2, "sam_cli": 10000, "sam_odonto": 7500, "mult_cac": 1.3}),
]
for i, (label, ov) in enumerate(STRESS):
    cases[f"st_{i}"] = dict(ov, cenario=2)

TORNADO = [
    ("Churn (todos os planos, ×0,67 a ×1,5)", "mult_churn", 1.5, 0.67),
    ("CAC (×1,5 a ×0,67)", "mult_cac", 1.5, 0.67),
    ("Clínicas por especialista (×0,67 a ×1,5)", "mult_esp", 0.67, 1.5),
    ("Preços (×0,8 a ×1,2)", "mult_preco", 0.8, 1.2),
    ("Receita de expansão (0% a 1% ao mês)", "exp_taxa", 0.0, 0.01),
    ("Custo do especialista (R$ 14 mil a R$ 10 mil)", "custo_esp", 14000, 10000),
    ("Custo de implantação (R$ 3.000 a R$ 1.000)", "impl", 3000, 1000),
    ("Indicação (0,4% a 1,2% da base por mês)", "indic", 0.004, 0.012),
]
for j, (label, key, lo, hi) in enumerate(TORNADO):
    cases[f"t_{j}_lo"] = {"cenario": 2, key: lo}
    cases[f"t_{j}_hi"] = {"cenario": 2, key: hi}

# condições mínimas: grade de multiplicadores, uma variável por vez, no base e no conservador
GRID = {
    "mult_esp": [0.5, 0.75, 1.0, 1.25, 1.5, 1.75, 2.0, 2.5, 3.0],
    "mult_churn": [0.5, 0.75, 1.0, 1.25, 1.5, 1.75, 2.0, 2.5],
    "mult_cac": [0.5, 0.75, 1.0, 1.25, 1.5, 2.0, 2.5, 3.0],
    "mult_preco": [0.7, 0.8, 0.9, 1.0, 1.1, 1.2, 1.3, 1.5],
}
for cen in (1, 2):
    for key, vals in GRID.items():
        for v in vals:
            cases[f"g{cen}_{key}_{v}"] = {"cenario": cen, key: v}

paths = {name: build(name, ov) for name, ov in cases.items()}
calc = recalc_all(list(paths.values()))
R = {name: read(calc[p]) for name, p in paths.items()}
json.dump(R, open(os.path.join(HERE, "resultados_clinicas.json"), "w"), indent=1, ensure_ascii=False)


def threshold(cen, key, metric, crit, higher_is_better):
    """Menor (ou maior) multiplicador da grade que ainda passa no critério, com interpolação linear."""
    pts = [(v, num(R[f"g{cen}_{key}_{v}"][metric])) for v in GRID[key]]
    ok = [(v, y) for v, y in pts if y is not None]
    for (v0, y0), (v1, y1) in zip(ok, ok[1:]):
        if (y0 - crit) * (y1 - crit) <= 0 and y0 != y1:
            return v0 + (crit - y0) * (v1 - v0) / (y1 - y0)
    good = [v for v, y in ok if (y >= crit) == higher_is_better]
    return "passa em toda a grade" if len(good) == len(ok) else "não passa na grade"


MIN = {}
for cen in (1, 2):
    MIN[cen] = {
        "mb_esp": threshold(cen, "mult_esp", "mb5", 0.70, True),
        "mb_preco": threshold(cen, "mult_preco", "mb5", 0.70, True),
        "ltv_churn": threshold(cen, "mult_churn", "ltvcac", 3, True),
        "ltv_cac": threshold(cen, "mult_cac", "ltvcac", 3, True),
        "pb_cac": threshold(cen, "mult_cac", "payback", 12, False),
    }
json.dump({"resultados": R, "minimos": MIN}, open(os.path.join(HERE, "resultados_clinicas.json"), "w"), indent=1, ensure_ascii=False)

# ------------------------------------------------------------------ grava no xlsx final
subprocess.run(["python3", BUILD, FINAL], check=True, capture_output=True)
wb = openpyxl.load_workbook(FINAL)
BOLD = Font(name="Arial", bold=True)
TITLE = Font(name="Arial", bold=True, size=14, color="1F2440")
H2 = Font(name="Arial", bold=True, color="FFFFFF")
HFILL = PatternFill("solid", fgColor="1F2440")
FMT = {"arr5": 'R$ #,##0', "arr7": 'R$ #,##0', "cap": 'R$ #,##0', "cac5": 'R$ #,##0', "mb5": '0.0%', "pen50": '0.0%', "ltvcac": '0.0"x"', "payback": '0.0', "eq": '0', "m50": '0', "cli5": '#,##0'}


def hdr(ws, r, labels):
    for i, lab in enumerate(labels):
        c = ws.cell(row=r, column=1 + i, value=lab)
        c.font = H2
        c.fill = HFILL
        c.alignment = Alignment(wrap_text=True, horizontal="left" if i == 0 else "center")


C = wb.create_sheet("Cenarios")
C.column_dimensions["A"].width = 56
for col in range(2, 16):
    C.column_dimensions[openpyxl.utils.get_column_letter(col)].width = 15
C.cell(1, 1, "Cenários e testes de estresse (valores gerados por sensibilidade_clinicas.py)").font = TITLE
r = 3
hdr(C, r, ["Indicador", "Conservador", "Base", "Otimista"]); r += 1
for key, label in OUTPUTS:
    C.cell(r, 1, label)
    for i, c in enumerate((1, 2, 3)):
        C.cell(r, 2 + i, R[f"cen_{c}"][key]).number_format = FMT[key]
    r += 1
r += 1
C.cell(r, 1, "Testes de estresse (cenário base, uma mudança por vez, salvo indicação)").font = BOLD
r += 1
hdr(C, r, ["Teste"] + [label for _, label in OUTPUTS]); r += 1
C.row_dimensions[r - 1].height = 60
for i, (label, _) in enumerate(STRESS):
    C.cell(r, 1, label)
    for j, (key, _) in enumerate(OUTPUTS):
        C.cell(r, 2 + j, R[f"st_{i}"][key]).number_format = FMT[key]
    r += 1
r += 1
C.cell(r, 1, "Condições mínimas: multiplicador sobre a premissa do cenário para passar no critério (uma variável por vez)").font = BOLD
r += 1
hdr(C, r, ["Critério e variável", "No cenário base", "No conservador"]); r += 1
for key, label in (("mb_esp", "Margem de 70%: clínicas por especialista (× premissa)"), ("mb_preco", "Margem de 70%: preços (× premissa)"),
                   ("ltv_churn", "LTV/CAC de 3: churn máximo (× premissa)"), ("ltv_cac", "LTV/CAC de 3: CAC máximo (× premissa)"),
                   ("pb_cac", "Recuperação em 12 meses: CAC máximo (× premissa)")):
    C.cell(r, 1, label)
    for i, cen in enumerate((2, 1)):
        v = MIN[cen][key]
        C.cell(r, 2 + i, round(v, 2) if isinstance(v, float) else v)
    r += 1

T = wb.create_sheet("Sensibilidade")
T.column_dimensions["A"].width = 46
for col in range(2, 10):
    T.column_dimensions[openpyxl.utils.get_column_letter(col)].width = 17
T.cell(1, 1, "Sensibilidade no cenário base (valores gerados por script)").font = TITLE
r = 3
hdr(T, r, ["Premissa (ordenada pelo impacto)", "Valor pessimista", "Valor otimista", "ARR ano 7 (pess.)", "ARR ano 7 (ot.)", "Capital (pess.)", "Capital (ot.)", "Margem ano 5 (pess. / ot.)", "LTV/CAC (pess. / ot.)"]); r += 1
rows = []
for j, (label, key, lo, hi) in enumerate(TORNADO):
    a, b = R[f"t_{j}_lo"], R[f"t_{j}_hi"]
    rows.append((abs((b["cap"] or 0) - (a["cap"] or 0)) + abs((b["arr7"] or 0) - (a["arr7"] or 0)) / 10, label, lo, hi, a, b))
for amp, label, lo, hi, a, b in sorted(rows, reverse=True):
    T.cell(r, 1, label)
    T.cell(r, 2, lo)
    T.cell(r, 3, hi)
    for col, v, f in ((4, a["arr7"], FMT["arr7"]), (5, b["arr7"], FMT["arr7"]), (6, a["cap"], FMT["cap"]), (7, b["cap"], FMT["cap"])):
        T.cell(r, col, v).number_format = f
    T.cell(r, 8, f"{a['mb5']:.1%} / {b['mb5']:.1%}")
    T.cell(r, 9, f"{a['ltvcac']:.1f}x / {b['ltvcac']:.1f}x")
    r += 1
wb.save(FINAL)

for c in (1, 2, 3):
    print("cen", c, R[f"cen_{c}"], passes(R[f"cen_{c}"]))
for i, (label, _) in enumerate(STRESS):
    print("st", label, R[f"st_{i}"], passes(R[f"st_{i}"]))
for j, (label, key, lo, hi) in enumerate(TORNADO):
    a, b = R[f"t_{j}_lo"], R[f"t_{j}_hi"]
    print("tor", label, f"arr7 {a['arr7']:,.0f}..{b['arr7']:,.0f} cap {a['cap']:,.0f}..{b['cap']:,.0f} mb {a['mb5']:.3f}..{b['mb5']:.3f} ltv {a['ltvcac']:.1f}..{b['ltvcac']:.1f}")
print("MIN", MIN)
