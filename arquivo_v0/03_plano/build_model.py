"""Gera o modelo financeiro da Jalecon (xlsx com fórmulas)."""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.comments import Comment
from openpyxl.chart import LineChart, BarChart, Reference

OUT = "/home/user/jalecon/arquivo_v0/03_plano/Jalecon_Modelo_Financeiro.xlsx"
FONT = "Arial"
BLUE = Font(name=FONT, color="0000FF")
BLACK = Font(name=FONT, color="000000")
GREEN = Font(name=FONT, color="008000")
BOLD = Font(name=FONT, bold=True)
TITLE = Font(name=FONT, bold=True, size=14, color="1F2440")
H2 = Font(name=FONT, bold=True, size=11, color="FFFFFF")
HFILL = PatternFill("solid", fgColor="1F2440")
SFILL = PatternFill("solid", fgColor="E9EBF2")
YFILL = PatternFill("solid", fgColor="FFFF00")
THIN = Border(bottom=Side(style="thin", color="D7DAE5"))
BRL = 'R$ #,##0;(R$ #,##0);-'
BRL2 = 'R$ #,##0.00;(R$ #,##0.00);-'
NUM = '#,##0;(#,##0);-'
NUM1 = '#,##0.0;(#,##0.0);-'
PCT = '0.0%;(0.0%);-'
MULT = '0.0"x"'

wb = Workbook()
A = {}  # endereços nomeados


def cell(ws, ref, value, font=BLACK, fmt=None, fill=None, comment=None):
    c = ws[ref]
    c.value = value
    c.font = font
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = fill
    if comment:
        c.comment = Comment(comment, "Jalecon")
    return c


def header(ws, row, labels, start_col=1):
    for i, lab in enumerate(labels):
        c = ws.cell(row=row, column=start_col + i, value=lab)
        c.font = H2
        c.fill = HFILL
        c.alignment = Alignment(horizontal="center" if i else "left")


def section(ws, row, text, width=8):
    for col in range(1, width + 1):
        ws.cell(row=row, column=col).fill = SFILL
    c = ws.cell(row=row, column=1, value=text)
    c.font = BOLD


# ------------------------------------------------------------------ Premissas
P = wb.active
P.title = "Premissas"
P.column_dimensions["A"].width = 52
for col in "BCDEFGH":
    P.column_dimensions[col].width = 16
P.column_dimensions["I"].width = 60
cell(P, "A1", "Jalecon: premissas do modelo financeiro", TITLE)
cell(P, "A2", "Legenda: texto azul = premissa editável · preto = fórmula · verde = vínculo com outra aba · fundo amarelo = premissa-chave a validar. Valores em R$ nominais, mês 1 = jan/2027.", Font(name=FONT, italic=True, size=9))

r = 4
section(P, r, "Geral", 9); r += 1
gen = [
    ("cambio", "Câmbio (R$ por US$)", 5.40, '0.00', "Premissa do grupo; ajustar à cotação do dia da banca."),
    ("simples", "Carga tributária no Simples Nacional (% da receita)", 0.11, PCT, "Alíquota efetiva média estimada para SaaS no Anexo III/V; confirmar com contador."),
    ("presumido", "Carga tributária após sair do Simples (% da receita)", 0.1633, PCT, "Aproximação de Lucro Presumido (PIS/COFINS 3,65% + ISS ~2,9% + IRPJ/CSLL presumidos ~9,8%); confirmar com contador."),
    ("limite", "Limite anual do Simples Nacional (R$)", 4800000, BRL, "Lei Complementar 123/2006."),
    ("pagto", "Taxa de meio de pagamento (% da receita)", 0.035, PCT, "Média de gateways de cartão/Pix recorrente."),
    ("enc_clt", "Encargos e benefícios CLT (% do salário)", 0.70, PCT, "Média de mercado para CLT no Brasil."),
    ("enc_pl", "Encargos sobre pró-labore (% do pró-labore)", 0.20, PCT, "INSS patronal no pró-labore (fora do Simples Anexo III)."),
    ("pct_anual", "% de clientes no plano anual", 0.30, PCT, "Premissa; validar na pré-venda."),
    ("desc_anual", "Desconto do plano anual", 0.15, PCT, "Premissa de política comercial."),
    ("preseed", "Aporte pré-seed no mês 1 (R$)", 3300000, BRL, "Valor a decidir após o modelo; ver aba Resumo (necessidade de caixa)."),
]
for key, lab, val, fmt, note in gen:
    cell(P, f"A{r}", lab)
    cell(P, f"B{r}", val, BLUE, fmt, YFILL if key in ("preseed",) else None)
    cell(P, f"I{r}", note, Font(name=FONT, size=9, color="4A4F6E"))
    A[key] = f"Premissas!$B${r}"
    r += 1

r += 1
section(P, r, "Planos", 9); r += 1
plans = ["Essencial", "Consultório", "Autoridade", "Clínica"]
header(P, r, ["Item", *plans, "", "", "", "Nota"]); r += 1
plan_rows = [
    ("preco", "Preço mensal (R$)", [297, 697, 1497, 3900], BRL, "Ancorado entre geradores de post (R$ 50 a 600) e agências (R$ 1,5 mil a 12 mil). Validar com Van Westendorp."),
    ("setup", "Taxa de implantação média por novo cliente (R$)", [0, 490, 990, 2900], BRL, "Marca + site + diagnóstico completo; isenta no plano anual (já refletido na média)."),
    ("churn", "Churn mensal no ano 1 (cenário base)", [0.05, 0.03, 0.025, 0.02], PCT, "SaaS para pequenos negócios tem mediana perto de 3,5%/mês; planos maiores retêm mais."),
    ("ia_usd", "Custo de IA por conta/mês (US$)", [4, 9, 18, 45], '0.00', "Estimativa: ~0,6 a 1,5 milhão de tokens/mês por conta com Claude Sonnet 5.5 (US$ 2/10 por milhão) e cache; inclui imagens."),
    ("infra", "Infra, APIs e hospedagem por conta/mês (R$)", [8, 12, 18, 40], BRL, "Hospedagem de site, banco de dados, APIs de publicação e WhatsApp."),
    ("cap_gestor", "Contas por gestor de conta humano", [0, 0, 40, 20], NUM, "Só planos com revisão humana dedicada. 0 = sem gestor."),
]
for key, lab, vals, fmt, note in plan_rows:
    cell(P, f"A{r}", lab)
    for j, v in enumerate(vals):
        cell(P, f"{L(2+j)}{r}", v, BLUE, fmt, YFILL if key in ("preco", "churn") else None)
    cell(P, f"I{r}", note, Font(name=FONT, size=9, color="4A4F6E"))
    A[key] = r
    r += 1
cell(P, f"A{r}", "Salário mensal do gestor de conta (R$)")
cell(P, f"B{r}", 7000, BLUE, BRL)
A["sal_gestor"] = f"Premissas!$B${r}"; r += 1
cell(P, f"A{r}", "Melhoria anual do churn (redução relativa por ano)")
cell(P, f"B{r}", 0.08, BLUE, PCT)
cell(P, f"I{r}", "Produto e sucesso do cliente amadurecem; aplicado a partir do ano 2.", Font(name=FONT, size=9, color="4A4F6E"))
A["melhoria"] = f"Premissas!$B${r}"; r += 1

r += 1
section(P, r, "Mix de novos clientes por plano", 9); r += 1
header(P, r, ["Ano", *plans, "Soma"]); r += 1
mix = [[0.40, 0.50, 0.10, 0.00], [0.35, 0.50, 0.13, 0.02], [0.30, 0.50, 0.15, 0.05], [0.28, 0.50, 0.16, 0.06], [0.25, 0.50, 0.18, 0.07]]
A["mix_first"] = r
for y, row in enumerate(mix):
    cell(P, f"A{r}", f"Ano {y+1}")
    for j, v in enumerate(row):
        cell(P, f"{L(2+j)}{r}", v, BLUE, PCT)
    cell(P, f"F{r}", f"=SUM(B{r}:E{r})", BLACK, PCT)
    r += 1

r += 1
section(P, r, "Aquisição e receitas adicionais (cenário base)", 9); r += 1
header(P, r, ["Item", "Ano 1", "Ano 2", "Ano 3", "Ano 4", "Ano 5", "", "Nota"]); r += 1
acq = [
    ("mkt", "Mídia paga e inbound: orçamento mensal (R$)", [25000, 70000, 150000, 240000, 320000], BRL, "Meta, Google, conteúdo, ferramentas de captação."),
    ("parc", "Parcerias: novos clientes por mês", [4, 12, 25, 40, 55], NUM, "Residências, sociedades de especialidade, contadores médicos, coworkings de consultório."),
    ("addon", "Serviços adicionais (% da receita recorrente)", [0.05, 0.08, 0.11, 0.13, 0.15], PCT, "Check-up de conformidade avulso, gestão de tráfego, ensaio de fotos (comissão), domínio."),
]
for key, lab, vals, fmt, note in acq:
    cell(P, f"A{r}", lab)
    for j, v in enumerate(vals):
        cell(P, f"{L(2+j)}{r}", v, BLUE, fmt)
    cell(P, f"H{r}", note, Font(name=FONT, size=9, color="4A4F6E"))
    A[key] = r
    r += 1
cell(P, f"A{r}", "CAC de mídia paga no ano 1 (R$ por cliente)")
cell(P, f"B{r}", 1800, BLUE, BRL, YFILL)
cell(P, f"H{r}", "Premissa-chave: SaaS para pequenos negócios com venda assistida. Validar com testes de campanha.", Font(name=FONT, size=9, color="4A4F6E"))
cac1 = r; r += 1
cell(P, f"A{r}", "Queda anual do CAC (marca, indicação, conteúdo)")
cell(P, f"B{r}", 0.07, BLUE, PCT)
qcac = r; r += 1
cell(P, f"A{r}", "CAC de mídia paga por ano (R$)")
for j in range(5):
    if j == 0:
        cell(P, f"{L(2+j)}{r}", f"=$B${cac1}", BLACK, BRL)
    else:
        cell(P, f"{L(2+j)}{r}", f"={L(1+j)}{r}*(1-$B${qcac})", BLACK, BRL)
A["cac"] = r; r += 1
cell(P, f"A{r}", "Indicação: novos clientes por mês (% da base ativa)")
cell(P, f"B{r}", 0.01, BLUE, PCT)
A["indic"] = f"Premissas!$B${r}"; r += 1

r += 1
section(P, r, "Cenários", 9); r += 1
header(P, r, ["Alavanca", "Conservador", "Base", "Otimista", "", "", "", "Nota"]); r += 1
scen = [
    ("m_cac", "Multiplicador do CAC", [1.3, 1.0, 0.8], MULT, "CAC 30% maior no conservador, 20% menor no otimista."),
    ("m_churn", "Multiplicador do churn", [1.3, 1.0, 0.8], MULT, ""),
    ("m_org", "Multiplicador de parcerias e indicação", [0.7, 1.0, 1.3], MULT, ""),
    ("m_addon", "Multiplicador de serviços adicionais", [0.7, 1.0, 1.2], MULT, ""),
    ("m_mkt", "Multiplicador do orçamento de mídia", [0.8, 1.0, 1.15], MULT, "No conservador a empresa gasta menos por ter menos caixa."),
    ("seed", "Rodada seed (R$)", [5000000, 8000000, 10000000], BRL, "Premissa de captação futura; não é compromisso."),
    ("seed_mes", "Mês da rodada seed", [20, 16, 14], NUM, ""),
]
for key, lab, vals, fmt, note in scen:
    cell(P, f"A{r}", lab)
    for j, v in enumerate(vals):
        cell(P, f"{L(2+j)}{r}", v, BLUE, fmt)
    cell(P, f"H{r}", note, Font(name=FONT, size=9, color="4A4F6E"))
    A[key] = r
    r += 1

r += 1
section(P, r, "Equipe (custo mensal por pessoa e quadro por ano)", 9); r += 1
header(P, r, ["Função", "Salário (R$)", "Encargos", "Ano 1", "Ano 2", "Ano 3", "Ano 4", "Ano 5", "Nota"]); r += 1
team = [
    ("Engenharia de software", 14000, "clt", [2, 4, 7, 10, 13], "Plataforma, integrações (Meta, Google, WhatsApp, agendas)."),
    ("Engenharia de IA e dados", 18000, "clt", [1, 2, 3, 4, 5], "Agentes, Guardião, avaliação de qualidade, dados de conversão."),
    ("Produto e design", 12000, "clt", [1, 2, 3, 4, 5], "Ano 1: designer de produto; gestão de produto com a fundadora (Mariana)."),
    ("Marketing e growth", 10000, "clt", [0, 2, 3, 5, 6], "Ano 1 conduzido pelo fundador de marketing (Vitor)."),
    ("Vendas (inside sales)", 7000, "clt", [1, 3, 6, 9, 12], "Salário fixo; comissão incluída no CAC."),
    ("Sucesso do cliente e suporte", 6500, "clt", [1, 2, 4, 7, 10], "Onboarding e suporte geral (gestores de conta dedicados estão no custo do serviço)."),
    ("Médicos revisores de conteúdo (PJ)", 9000, "pj", [1, 1, 2, 3, 4], "Revisão científica por especialidade."),
    ("Financeiro e administrativo", 8000, "clt", [0, 1, 2, 3, 4], ""),
]
team_first = r
for lab, sal, kind, hc, note in team:
    cell(P, f"A{r}", lab)
    cell(P, f"B{r}", sal, BLUE, BRL)
    cell(P, f"C{r}", f"={A['enc_clt'].split('!')[1]}" if kind == "clt" else 0, BLACK if kind == "clt" else BLUE, PCT)
    for j, v in enumerate(hc):
        cell(P, f"{L(4+j)}{r}", v, BLUE, NUM)
    cell(P, f"I{r}", note, Font(name=FONT, size=9, color="4A4F6E"))
    r += 1
team_last = r - 1
cell(P, f"A{r}", "Fundadores (quantidade)")
cell(P, f"C{r}", f"={A['enc_pl'].split('!')[1]}", BLACK, PCT)
for j, v in enumerate([5, 5, 5, 5, 5]):
    cell(P, f"{L(4+j)}{r}", v, BLUE, NUM)
cell(P, f"I{r}", "Mariana, Vitor, Fabiana, Manuella e Alessandra em tempo integral; Daniela como estagiária/part-time a definir.", Font(name=FONT, size=9, color="4A4F6E"))
fr = r; r += 1
cell(P, f"A{r}", "Pró-labore mensal por fundador (R$)")
for j, v in enumerate([5000, 10000, 15000, 20000, 22000]):
    cell(P, f"{L(4+j)}{r}", v, BLUE, BRL)
cell(P, f"I{r}", "Baixo no ano 1 para preservar caixa; sobe com a receita.", Font(name=FONT, size=9, color="4A4F6E"))
pl = r; r += 1
cell(P, f"A{r}", "Custo mensal total de pessoal (R$)", BOLD)
for j in range(5):
    col = L(4 + j)
    cell(P, f"{col}{r}", f"=SUMPRODUCT($B${team_first}:$B${team_last},1+$C${team_first}:$C${team_last},{col}{team_first}:{col}{team_last})+{col}{fr}*{col}{pl}*(1+$C${fr})", BLACK, BRL)
A["pessoal"] = r; r += 1
cell(P, f"A{r}", "Quadro total (pessoas, sem gestores de conta)")
for j in range(5):
    col = L(4 + j)
    cell(P, f"{col}{r}", f"=SUM({col}{team_first}:{col}{team_last})+{col}{fr}", BLACK, NUM)
A["hc"] = r; r += 1

r += 1
section(P, r, "Outras despesas mensais (R$)", 9); r += 1
header(P, r, ["Item", "Ano 1", "Ano 2", "Ano 3", "Ano 4", "Ano 5", "", "Nota"]); r += 1
other = [
    ("juridico", "Jurídico e curadoria normativa (advogado de direito médico)", [5000, 8000, 12000, 18000, 25000], "Mantém a biblioteca do Guardião atualizada."),
    ("eventos", "Marca: eventos, congressos e conteúdo institucional", [2000, 8000, 15000, 25000, 35000], ""),
    ("ferramentas", "Ferramentas e software internos", [3000, 8000, 15000, 25000, 35000], ""),
    ("escritorio", "Escritório / coworking", [2000, 6000, 15000, 25000, 35000], ""),
    ("ga", "Contabilidade, seguros e outras despesas", [3000, 5000, 9000, 14000, 20000], ""),
]
for key, lab, vals, note in other:
    cell(P, f"A{r}", lab)
    for j, v in enumerate(vals):
        cell(P, f"{L(2+j)}{r}", v, BLUE, BRL)
    cell(P, f"H{r}", note, Font(name=FONT, size=9, color="4A4F6E"))
    A[key] = r
    r += 1
P.freeze_panes = "B4"

# ------------------------------------------------------------------ Modelos mensais
MONTHS = 60
FC = 3  # primeira coluna de mês (C)
LC = FC + MONTHS - 1
lastL = L(LC)
scen_names = {1: "Conservador", 2: "Base", 3: "Otimista"}
MROWS = {}


def prem_year(rowno, col):
    """Valor anual da Premissas (colunas B..F = anos 1..5)."""
    return f"INDEX(Premissas!$B${rowno}:$F${rowno},{col}$4)"


def prem_scen(key):
    rr = A[key]
    return f"INDEX(Premissas!$B${rr}:$D${rr},$B$2)"


def build_model(idx):
    ws = wb.create_sheet(f"Modelo_{scen_names[idx]}")
    ws.column_dimensions["A"].width = 46
    ws.column_dimensions["B"].width = 12
    for c in range(FC, LC + 1):
        ws.column_dimensions[L(c)].width = 12
    cell(ws, "A1", f"Jalecon: projeção mensal, cenário {scen_names[idx].lower()}", TITLE)
    cell(ws, "A2", "Cenário (1 = conservador, 2 = base, 3 = otimista)")
    cell(ws, "B2", idx, BLUE)
    cell(ws, "A3", "Mês")
    cell(ws, "A4", "Ano")
    for m in range(1, MONTHS + 1):
        col = L(FC + m - 1)
        cell(ws, f"{col}3", m, BLUE if m == 1 else BLACK)
        if m > 1:
            ws[f"{col}3"].value = f"={L(FC+m-2)}3+1"
        cell(ws, f"{col}4", f"=ROUNDUP({col}3/12,0)")
        ws[f"{col}3"].font = BOLD
    rows = {}
    rr = [6]

    def row(key, label, fn, fmt=NUM, bold=False):
        n = rr[0]
        rows[key] = n
        cell(ws, f"A{n}", label, BOLD if bold else BLACK)
        for m in range(1, MONTHS + 1):
            col = L(FC + m - 1)
            prev = L(FC + m - 2) if m > 1 else None
            f = fn(col, prev, m)
            c = cell(ws, f"{col}{n}", f, GREEN if "Premissas!" in str(f) else BLACK, fmt)
            if bold:
                c.font = Font(name=FONT, bold=True, color=c.font.color.rgb if c.font.color else "000000")
        rr[0] += 1
        return n

    def sec(label):
        section(ws, rr[0], label, 2)
        rr[0] += 1

    sec("Aquisição")
    row("mkt", "Orçamento de mídia paga (R$)", lambda c, p, m: f"={prem_year(A['mkt'], c)}*{prem_scen('m_mkt')}", BRL)
    row("cac", "CAC de mídia paga (R$)", lambda c, p, m: f"={prem_year(A['cac'], c)}*{prem_scen('m_cac')}", BRL)
    row("n_pago", "Novos clientes: mídia paga", lambda c, p, m: f"=IFERROR({c}{rows['mkt']}/{c}{rows['cac']},0)", NUM1)
    row("n_parc", "Novos clientes: parcerias", lambda c, p, m: f"={prem_year(A['parc'], c)}*{prem_scen('m_org')}", NUM1)
    row("n_ind", "Novos clientes: indicação", lambda c, p, m: "=0" if m == 1 else f"={p}{{ativos}}*{A['indic']}*{prem_scen('m_org')}", NUM1)
    row("n_tot", "Novos clientes: total", lambda c, p, m: f"={c}{rows['n_pago']}+{c}{rows['n_parc']}+{c}{rows['n_ind']}", NUM1, True)

    sec("Clientes por plano")
    for j, pn in enumerate(plans):
        colP = L(2 + j)
        row(f"nov{j}", f"Novos: {pn}", lambda c, p, m, j=j, colP=colP: f"={c}{rows['n_tot']}*INDEX(Premissas!${colP}${A['mix_first']}:${colP}${A['mix_first']+4},{c}$4)", NUM1)
    for j, pn in enumerate(plans):
        colP = L(2 + j)
        row(f"chu{j}", f"Cancelamentos: {pn}", lambda c, p, m, j=j, colP=colP: "=0" if m == 1 else f"={p}{{at{j}}}*Premissas!${colP}${A['churn']}*{prem_scen('m_churn')}*(1-{A['melhoria']})^({c}$4-1)", NUM1)
    for j, pn in enumerate(plans):
        row(f"at{j}", f"Ativos: {pn}", lambda c, p, m, j=j: f"={c}{{nov{j}}}-{c}{{chu{j}}}" if m == 1 else f"={p}{{at{j}}}+{c}{{nov{j}}}-{c}{{chu{j}}}", NUM)
    row("ativos", "Clientes ativos: total", lambda c, p, m: f"=SUM({c}{{at0}}:{c}{{at3}})", NUM, True)

    sec("Receita (R$)")
    fator = f"(1-{A['pct_anual']}*{A['desc_anual']})"
    for j, pn in enumerate(plans):
        colP = L(2 + j)
        row(f"mrr{j}", f"Receita recorrente: {pn}", lambda c, p, m, j=j, colP=colP: f"={c}{{at{j}}}*Premissas!${colP}${A['preco']}*{fator}", BRL)
    row("mrr", "Receita recorrente mensal (MRR)", lambda c, p, m: f"=SUM({c}{{mrr0}}:{c}{{mrr3}})", BRL, True)
    row("setup", "Taxas de implantação", lambda c, p, m: "=" + "+".join(f"{c}{{nov{j}}}*Premissas!${L(2+j)}${A['setup']}" for j in range(4)), BRL)
    row("addon", "Serviços adicionais", lambda c, p, m: f"={c}{rows['mrr']}*{prem_year(A['addon'], c)}*{prem_scen('m_addon')}", BRL)
    row("rec", "Receita bruta", lambda c, p, m: f"={c}{rows['mrr']}+{c}{rows['setup']}+{c}{rows['addon']}", BRL, True)
    row("imp", "(-) Impostos sobre receita", lambda c, p, m: f"=-{c}{rows['rec']}*IF(SUMIFS(${L(FC)}${rows['rec']}:${lastL}${rows['rec']},${L(FC)}$4:${lastL}$4,{c}$4-1)>{A['limite']},{A['presumido']},{A['simples']})", BRL)
    row("rliq", "Receita líquida", lambda c, p, m: f"={c}{rows['rec']}+{c}{rows['imp']}", BRL, True)

    sec("Custo do serviço (R$)")
    row("c_ia", "IA (modelos de linguagem e imagem)", lambda c, p, m: "=-(" + "+".join(f"{c}{{at{j}}}*Premissas!${L(2+j)}${A['ia_usd']}" for j in range(4)) + f")*{A['cambio']}", BRL)
    row("c_infra", "Infraestrutura e APIs", lambda c, p, m: "=-(" + "+".join(f"{c}{{at{j}}}*Premissas!${L(2+j)}${A['infra']}" for j in range(4)) + ")", BRL)
    row("c_gestor", "Gestores de conta (revisão humana)", lambda c, p, m: "=-(" + "+".join(f"IFERROR({c}{{at{j}}}/Premissas!${L(2+j)}${A['cap_gestor']},0)" for j in range(4)) + f")*{A['sal_gestor']}*(1+{A['enc_clt']})", BRL)
    row("c_pag", "Meio de pagamento", lambda c, p, m: f"=-{c}{rows['rec']}*{A['pagto']}", BRL)
    row("cogs", "Custo do serviço total", lambda c, p, m: f"=SUM({c}{rows['c_ia']}:{c}{rows['c_pag']})", BRL, True)
    row("lb", "Lucro bruto", lambda c, p, m: f"={c}{rows['rliq']}+{c}{rows['cogs']}", BRL, True)
    row("mb", "Margem bruta (% da receita líquida)", lambda c, p, m: f"=IFERROR({c}{rows['lb']}/{c}{rows['rliq']},0)", PCT)

    sec("Despesas operacionais (R$)")
    row("o_pes", "Pessoal", lambda c, p, m: f"=-INDEX(Premissas!$D${A['pessoal']}:$H${A['pessoal']},{c}$4)", BRL)
    row("o_mkt", "Mídia paga e inbound", lambda c, p, m: f"=-{c}{rows['mkt']}", BRL)
    row("o_com", "Comissões de vendas (10% do primeiro mês dos novos)", lambda c, p, m: "=-0.1*(" + "+".join(f"{c}{{nov{j}}}*Premissas!${L(2+j)}${A['preco']}" for j in range(4)) + ")", BRL)
    for key, lab in [("juridico", "Jurídico e curadoria normativa"), ("eventos", "Marca: eventos e congressos"), ("ferramentas", "Ferramentas internas"), ("escritorio", "Escritório"), ("ga", "Contabilidade, seguros e outros")]:
        row(f"o_{key}", lab, lambda c, p, m, key=key: f"=-{prem_year(A[key], c)}", BRL)
    row("opex", "Despesas operacionais totais", lambda c, p, m: f"=SUM({c}{rows['o_pes']}:{c}{rows['o_ga']})", BRL, True)
    row("ebitda", "EBITDA", lambda c, p, m: f"={c}{rows['lb']}+{c}{rows['opex']}", BRL, True)
    row("mebitda", "Margem EBITDA (% da receita líquida)", lambda c, p, m: f"=IFERROR({c}{rows['ebitda']}/{c}{rows['rliq']},0)", PCT)

    sec("Caixa (R$)")
    row("aporte", "Aportes de capital", lambda c, p, m: f"=IF({c}$3=1,{A['preseed']},0)+IF({c}$3={prem_scen('seed_mes')},{prem_scen('seed')},0)", BRL)
    row("caixa", "Saldo de caixa no fim do mês", lambda c, p, m: f"={c}{rows['aporte']}+{c}{rows['ebitda']}" if m == 1 else f"={p}{{caixa}}+{c}{rows['aporte']}+{c}{rows['ebitda']}", BRL, True)

    sec("Indicadores")
    row("arr", "ARR (MRR × 12)", lambda c, p, m: f"={c}{rows['mrr']}*12", BRL)
    row("arpa", "Receita recorrente média por cliente (ARPA)", lambda c, p, m: f"=IFERROR({c}{rows['mrr']}/{c}{rows['ativos']},0)", BRL)
    row("churn_t", "Churn mensal total (% da base do mês anterior)", lambda c, p, m: "=0" if m == 1 else f"=IFERROR(SUM({c}{{chu0}}:{c}{{chu3}})/{p}{rows['ativos']},0)", PCT)
    row("be_flag", "Mês com EBITDA positivo daí em diante (auxiliar)", lambda c, p, m: f"=IF(MIN({c}{rows['ebitda']}:${lastL}{rows['ebitda']})>0,{c}$3,999)", NUM)

    # resolver placeholders {key}
    for rown in range(6, rr[0]):
        for m in range(1, MONTHS + 1):
            c = ws.cell(row=rown, column=FC + m - 1)
            if isinstance(c.value, str) and "{" in c.value:
                v = c.value
                for k, n in rows.items():
                    v = v.replace("{" + k + "}", str(n))
                c.value = v
    ws.freeze_panes = "C5"
    MROWS[idx] = rows
    return ws


for i in (1, 2, 3):
    build_model(i)

# ------------------------------------------------------------------ Resumo anual
R = wb.create_sheet("Resumo", 1)
R.column_dimensions["A"].width = 46
for c in "BCDEFGH":
    R.column_dimensions[c].width = 16
cell(R, "A1", "Jalecon: resumo anual e comparação de cenários", TITLE)
cell(R, "A2", "Todos os valores vêm das abas Modelo_*. Para mudar premissas, edite a aba Premissas.", Font(name=FONT, italic=True, size=9))


def ysum(sheet, key, y):
    rows = MROWS[2] if sheet == "Modelo_Base" else MROWS[[k for k, v in scen_names.items() if f"Modelo_{v}" == sheet][0]]
    n = rows[key]
    return f"SUMIFS({sheet}!${L(FC)}${n}:${lastL}${n},{sheet}!${L(FC)}$4:${lastL}$4,{y})"


def yend(sheet, key, y):
    idx = [k for k, v in scen_names.items() if f"Modelo_{v}" == sheet][0]
    n = MROWS[idx][key]
    return f"INDEX({sheet}!${L(FC)}${n}:${lastL}${n},{y}*12)"


r = 4
section(R, r, "DRE anual: cenário base (R$)", 7); r += 1
header(R, r, ["Linha", "Ano 1 (2027)", "Ano 2 (2028)", "Ano 3 (2029)", "Ano 4 (2030)", "Ano 5 (2031)"]); r += 1
dre = [
    ("Clientes ativos no fim do ano", "ativos", "end", NUM),
    ("ARR no fim do ano", "arr", "end", BRL),
    ("Receita recorrente", "mrr", "sum", BRL),
    ("Taxas de implantação", "setup", "sum", BRL),
    ("Serviços adicionais", "addon", "sum", BRL),
    ("Receita bruta", "rec", "sum", BRL),
    ("(-) Impostos", "imp", "sum", BRL),
    ("Receita líquida", "rliq", "sum", BRL),
    ("(-) Custo do serviço", "cogs", "sum", BRL),
    ("Lucro bruto", "lb", "sum", BRL),
    ("Margem bruta", None, "mb", PCT),
    ("(-) Pessoal", "o_pes", "sum", BRL),
    ("(-) Mídia, comissões e marca", None, "mkt", BRL),
    ("(-) Demais despesas", None, "outras", BRL),
    ("EBITDA", "ebitda", "sum", BRL),
    ("Margem EBITDA", None, "me", PCT),
    ("Caixa no fim do ano", "caixa", "end", BRL),
]
dre_rows = {}
for lab, key, kind, fmt in dre:
    cell(R, f"A{r}", lab, BOLD if lab in ("Receita bruta", "Lucro bruto", "EBITDA", "Receita líquida") else BLACK)
    dre_rows[lab] = r
    for y in range(1, 6):
        col = L(1 + y)
        if kind == "sum":
            f = "=" + ysum("Modelo_Base", key, y)
        elif kind == "end":
            f = "=" + yend("Modelo_Base", key, y)
        elif kind == "mb":
            f = f"=IFERROR({col}{dre_rows['Lucro bruto']}/{col}{dre_rows['Receita líquida']},0)"
        elif kind == "me":
            f = f"=IFERROR({col}{dre_rows['EBITDA']}/{col}{dre_rows['Receita líquida']},0)"
        elif kind == "mkt":
            f = "=" + "+".join(ysum("Modelo_Base", k, y) for k in ("o_mkt", "o_com", "o_eventos"))
        elif kind == "outras":
            f = "=" + "+".join(ysum("Modelo_Base", k, y) for k in ("o_juridico", "o_ferramentas", "o_escritorio", "o_ga"))
        cell(R, f"{col}{r}", f, GREEN if kind in ("sum", "end", "mkt", "outras") else BLACK, fmt)
    r += 1
cell(R, f"A{r}", "Quadro de pessoas (sem gestores de conta)")
for y in range(1, 6):
    cell(R, f"{L(1+y)}{r}", f"=INDEX(Premissas!$D${A['hc']}:$H${A['hc']},{y})", GREEN, NUM)
r += 2

section(R, r, "Comparação de cenários", 7); r += 1
header(R, r, ["Indicador", "Conservador", "Base", "Otimista"]); r += 1
comp_first = r
comps = [
    ("Clientes ativos no fim do ano 3", lambda s: "=" + yend(s, "ativos", 3), NUM),
    ("Clientes ativos no fim do ano 5", lambda s: "=" + yend(s, "ativos", 5), NUM),
    ("ARR no fim do ano 3 (R$)", lambda s: "=" + yend(s, "arr", 3), BRL),
    ("ARR no fim do ano 5 (R$)", lambda s: "=" + yend(s, "arr", 5), BRL),
    ("Receita bruta do ano 5 (R$)", lambda s: "=" + ysum(s, "rec", 5), BRL),
    ("EBITDA do ano 5 (R$)", lambda s: "=" + ysum(s, "ebitda", 5), BRL),
    ("Mês do equilíbrio (EBITDA positivo daí em diante)", lambda s: f"=MIN({s}!${L(FC)}${MROWS[2]['be_flag']}:${lastL}${MROWS[2]['be_flag']})", NUM),
    ("Menor saldo de caixa no período (R$)", lambda s: f"=MIN({s}!${L(FC)}${MROWS[2]['caixa']}:${lastL}${MROWS[2]['caixa']})", BRL),
    ("Total captado no período (R$)", lambda s: f"=SUM({s}!${L(FC)}${MROWS[2]['aporte']}:${lastL}${MROWS[2]['aporte']})", BRL),
]
for lab, fn, fmt in comps:
    cell(R, f"A{r}", lab)
    for j, s in enumerate(["Modelo_Conservador", "Modelo_Base", "Modelo_Otimista"]):
        cell(R, f"{L(2+j)}{r}", fn(s), GREEN, fmt)
    r += 1
cell(R, f"A{r}", "Leitura: 999 em 'mês do equilíbrio' significa que o cenário não chega ao equilíbrio em 60 meses. Caixa negativo indica necessidade de captação adicional.", Font(name=FONT, italic=True, size=9))
r += 2

section(R, r, "Necessidade de caixa antes da rodada seed (cenário base)", 7); r += 1
seedm = f"INDEX(Premissas!$B${A['seed_mes']}:$D${A['seed_mes']},2)"
cell(R, f"A{r}", "Queima acumulada até o mês anterior ao seed (R$)")
eb = MROWS[2]["ebitda"]
cell(R, f"B{r}", f"=-SUMIFS(Modelo_Base!${L(FC)}${eb}:${lastL}${eb},Modelo_Base!${L(FC)}$3:${lastL}$3,\"<\"&{seedm})", GREEN, BRL)
burn_r = r; r += 1
cell(R, f"A{r}", "Reserva de segurança (meses de queima média)")
cell(R, f"B{r}", 3, BLUE, NUM)
res_r = r; r += 1
cell(R, f"A{r}", "Pré-seed recomendado (R$)", BOLD)
cell(R, f"B{r}", f"=ROUNDUP((B{burn_r}+B{burn_r}/({seedm}-1)*B{res_r})/100000,0)*100000", BLACK, BRL)
cell(R, f"C{r}", "Cobre a queima até o seed mais a reserva, arredondado para R$ 100 mil.", Font(name=FONT, size=9, color="4A4F6E"))
r += 2

# gráfico de clientes ativos por cenário (fim de ano)
section(R, r, "Clientes ativos no fim de cada ano por cenário (dados do gráfico)", 7); r += 1
header(R, r, ["Ano", "Conservador", "Base", "Otimista"]); r += 1
chart_first = r
for y in range(1, 6):
    cell(R, f"A{r}", f"Ano {y}")
    for j, s in enumerate(["Modelo_Conservador", "Modelo_Base", "Modelo_Otimista"]):
        cell(R, f"{L(2+j)}{r}", "=" + yend(s, "ativos", y), GREEN, NUM)
    r += 1
ch = LineChart()
ch.title = "Clientes ativos no fim do ano"
ch.y_axis.title = "Clientes"
ch.height = 8
ch.width = 18
data = Reference(R, min_col=2, min_row=chart_first - 1, max_col=4, max_row=chart_first + 4)
cats = Reference(R, min_col=1, min_row=chart_first, max_row=chart_first + 4)
ch.add_data(data, titles_from_data=True)
ch.set_categories(cats)
for s, colr in zip(ch.series, ["9C5A36", "1F2440", "3F6B57"]):
    s.graphicalProperties.line.solidFill = colr
    s.graphicalProperties.line.width = 28000
R.add_chart(ch, f"F{chart_first-1}")
r += 2
section(R, r, "Receita bruta e EBITDA por ano (cenário base, dados do gráfico)", 7); r += 1
header(R, r, ["Ano", "Receita bruta", "EBITDA"]); r += 1
c2 = r
for y in range(1, 6):
    cell(R, f"A{r}", f"Ano {y}")
    cell(R, f"B{r}", f"={L(1+y)}{dre_rows['Receita bruta']}", BLACK, BRL)
    cell(R, f"C{r}", f"={L(1+y)}{dre_rows['EBITDA']}", BLACK, BRL)
    r += 1
bc = BarChart()
bc.title = "Receita bruta e EBITDA (R$)"
bc.height = 8
bc.width = 18
bc.add_data(Reference(R, min_col=2, min_row=c2 - 1, max_col=3, max_row=c2 + 4), titles_from_data=True)
bc.set_categories(Reference(R, min_col=1, min_row=c2, max_row=c2 + 4))
for s, colr in zip(bc.series, ["1F2440", "E8A27C"]):
    s.graphicalProperties.solidFill = colr
R.add_chart(bc, f"F{c2-1}")
R.freeze_panes = "B4"

# ------------------------------------------------------------------ Unit economics
U = wb.create_sheet("Unit_Economics", 2)
U.column_dimensions["A"].width = 50
for c in "BCDEF":
    U.column_dimensions[c].width = 16
cell(U, "A1", "Jalecon: economia por cliente (cenário base, ano 1)", TITLE)
cell(U, "A2", "Margem de contribuição por conta e retorno do custo de aquisição.", Font(name=FONT, italic=True, size=9))
header(U, 4, ["Item", *plans])
ue = {}
items = [
    ("preco", "Preço efetivo mensal (R$, após desconto anual)", lambda j, cp: f"=Premissas!{cp}${A['preco']}*(1-{A['pct_anual']}*{A['desc_anual']})", BRL),
    ("imp", "(-) Impostos (Simples)", lambda j, cp: f"=-{{preco}}*{A['simples']}", BRL),
    ("ia", "(-) IA", lambda j, cp: f"=-Premissas!{cp}${A['ia_usd']}*{A['cambio']}", BRL),
    ("infra", "(-) Infra e APIs", lambda j, cp: f"=-Premissas!{cp}${A['infra']}", BRL),
    ("gestor", "(-) Gestor de conta (rateio)", lambda j, cp: f"=-IFERROR({A['sal_gestor']}*(1+{A['enc_clt']})/Premissas!{cp}${A['cap_gestor']},0)", BRL),
    ("pag", "(-) Meio de pagamento", lambda j, cp: f"=-{{preco}}*{A['pagto']}", BRL),
    ("mc", "Margem de contribuição mensal (R$)", lambda j, cp: "=SUM({preco}:{pag})", BRL),
    ("mcp", "Margem de contribuição (%)", lambda j, cp: "=IFERROR({mc}/{preco},0)", PCT),
    ("churn", "Churn mensal (ano 1, base)", lambda j, cp: f"=Premissas!{cp}${A['churn']}", PCT),
    ("vida", "Vida média do cliente (meses)", lambda j, cp: "=IFERROR(1/{churn},0)", NUM1),
    ("ltv", "LTV: margem acumulada na vida do cliente (R$)", lambda j, cp: "={mc}*{vida}", BRL),
    ("cac", "CAC de mídia paga, ano 1 (R$)", lambda j, cp: f"=Premissas!$B${A['cac']}", BRL),
    ("ltvcac", "LTV / CAC", lambda j, cp: "=IFERROR({ltv}/{cac},0)", MULT),
    ("payback", "Payback do CAC (meses)", lambda j, cp: "=IFERROR({cac}/{mc},0)", NUM1),
]
rr = 5
for key, lab, fn, fmt in items:
    ue[key] = rr
    rr += 1
for key, lab, fn, fmt in items:
    n = ue[key]
    cell(U, f"A{n}", lab, BOLD if key in ("mc", "ltvcac", "payback") else BLACK)
    for j in range(4):
        col = L(2 + j)
        cp = f"${col}"
        f = fn(j, cp)
        for k, v in ue.items():
            f = f.replace("{" + k + "}", f"{col}{v}")
        cell(U, f"{col}{n}", f, GREEN if "Premissas!" in f else BLACK, fmt)
n = rr + 1
cell(U, f"A{n}", "Referência de mercado: LTV/CAC acima de 3x e payback abaixo de 12 meses são considerados saudáveis para SaaS.", Font(name=FONT, italic=True, size=9))

# ------------------------------------------------------------------ Mercado
M = wb.create_sheet("Mercado", 3)
M.column_dimensions["A"].width = 58
M.column_dimensions["B"].width = 18
M.column_dimensions["C"].width = 70
cell(M, "A1", "Jalecon: tamanho de mercado (TAM, SAM, SOM)", TITLE)
header(M, 3, ["Item", "Valor", "Fonte / nota"])
mk = [
    ("med", "Médicos em atividade no Brasil (projeção 2025)", 635706, NUM, BLUE, "Demografia Médica no Brasil 2025 (FMUSP, AMB, Ministério da Saúde)."),
    ("auto", "% de médicos sem contrato formal (atuação autônoma/PJ)", 0.48, PCT, BLUE, "Demografia Médica 2025: subiu de 30% (2015) para 48% (2025)."),
    ("sudeste", "% dos especialistas no Sudeste", 0.554, PCT, BLUE, "Demografia Médica 2025."),
    ("ticket", "Receita anual média por conta (R$)", "=Unit_Economics!C5*12*(1+INDEX(Premissas!$B${a}:$F${a},3))".replace("{a}", str(A['addon'])), BRL, GREEN, "Plano Consultório + serviços adicionais do ano 3."),
    ("tam", "TAM: todos os médicos × receita anual por conta (R$)", "=B4*B7", BRL, BLACK, "Mercado total se todo médico assinasse."),
    ("sam_n", "SAM: médicos autônomos/PJ (quantidade)", "=B4*B5", NUM, BLACK, "Quem precisa gerar a própria demanda."),
    ("sam", "SAM (R$ por ano)", "=B9*B7", BRL, BLACK, ""),
    ("som_n", "SOM ano 5: clientes do cenário base", "=" + yend("Modelo_Base", "ativos", 5), NUM, GREEN, "Calculado pelo modelo."),
    ("som_p", "SOM como % do SAM", "=IFERROR(B11/B9,0)", PCT, BLACK, ""),
    ("som", "SOM (ARR do ano 5, cenário base, R$)", "=" + yend("Modelo_Base", "arr", 5), BRL, GREEN, ""),
    ("beach", "Primeiro segmento: autônomos no Sudeste (quantidade, aproximação)", "=B9*B6", NUM, BLACK, "Aproximação usando a participação do Sudeste nos especialistas."),
]
for i, (key, lab, val, fmt, font, note) in enumerate(mk):
    rr = 4 + i
    cell(M, f"A{rr}", lab)
    cell(M, f"B{rr}", val, font, fmt)
    cell(M, f"C{rr}", note, Font(name=FONT, size=9, color="4A4F6E"))
rr = 4 + len(mk) + 1
section(M, rr, "Expansão: outras profissões de saúde (estimativas a confirmar nos conselhos)", 3); rr += 1
exp = [
    ("Cirurgiões-dentistas", 400000, "Estimativa; confirmar no painel do CFO."),
    ("Psicólogos", 500000, "Estimativa; confirmar no CFP."),
    ("Nutricionistas", 150000, "Estimativa; o CFN cita 'mais de 110 mil' em nota antiga."),
    ("Fisioterapeutas", 300000, "Estimativa; confirmar no COFFITO."),
]
e0 = rr
for lab, v, note in exp:
    cell(M, f"A{rr}", lab)
    cell(M, f"B{rr}", v, BLUE, NUM, YFILL)
    cell(M, f"C{rr}", note, Font(name=FONT, size=9, color="4A4F6E"))
    rr += 1
cell(M, f"A{rr}", "Profissionais na expansão (total)", BOLD)
cell(M, f"B{rr}", f"=SUM(B{e0}:B{rr-1})", BLACK, NUM)
rr += 1
cell(M, f"A{rr}", "TAM ampliado (médicos + expansão, R$ por ano)", BOLD)
cell(M, f"B{rr}", f"=(B4+B{rr-1})*B7", BLACK, BRL)

# ------------------------------------------------------------------ Fontes
F = wb.create_sheet("Fontes_e_Notas")
F.column_dimensions["A"].width = 120
notes = [
    "Fontes de mercado",
    "Demografia Médica no Brasil 2025 (FMUSP, AMB, Ministério da Saúde): https://www.medway.com.br/conteudos/demografia-medica-no-brasil-2025-veja-os-principais-dados-do-novo-levantamento/",
    "Vínculos sem contrato formal 30% → 48%: https://www.medway.com.br/conteudos/como-esta-o-mercado-e-a-remuneracao-na-medicina/",
    "Preços de agências de marketing em saúde: https://farmacerto.com/quanto-custa-marketing-saude-brasil/ ; https://plusbusiness.com.br/blog/quanto-custa-agencia-de-marketing-digital/",
    "Preços de geradores de conteúdo médico com IA: https://www.pulmo.com.br/ ; https://instamed.temfavo.com/ ; https://www.doutorconteudo.com.br/ ; https://medcontent-ai.vercel.app/",
    "Preços de modelos de IA (Claude Sonnet 5.5: US$ 2 / US$ 10 por milhão de tokens de entrada/saída; Haiku 5.5: US$ 0,10 / 0,50; Batch com 50% de desconto): documentação da Anthropic, outubro/2026.",
    "",
    "Notas de método",
    "1. Clientes são calculados como valores esperados (podem ser fracionados); a exibição arredonda.",
    "2. Fluxo de caixa = EBITDA + aportes. Não considera capital de giro, capex relevante nem IR/CSLL além do que está na carga sobre receita.",
    "3. A troca de regime tributário usa a receita do ano anterior contra o limite do Simples.",
    "4. CAC de mídia paga inclui mídia e ferramentas de captação; comissões de vendas estão em linha própria; parcerias e indicação não têm CAC de mídia.",
    "5. Cenários alteram CAC, churn, parcerias/indicação, serviços adicionais, orçamento de mídia e a rodada seed. Todas as outras premissas são iguais.",
    "6. Tudo o que está em azul é premissa e deve ser revisado após as entrevistas de validação (Kit de Validação, documento 06).",
]
for i, t in enumerate(notes):
    c = F.cell(row=i + 1, column=1, value=t)
    c.font = BOLD if t in ("Fontes de mercado", "Notas de método") else Font(name=FONT)
    c.alignment = Alignment(wrap_text=True)

for ws in wb.worksheets:
    ws.sheet_view.showGridLines = False
wb.save(OUT)
print("salvo", OUT)
